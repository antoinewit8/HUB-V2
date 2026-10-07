"""
Page Streamlit : Itinéraires et Werbomont
Les itinéraires PTV chargement → déchargement de chaque ligne CIT, sur carte,
et pour chacune : passe-t-elle par Werbomont, et sinon quel détour, quelle renta ?

Fichiers : LISTES_MISSIONS (activités par dossier) + CA CIT (prix par dossier).
Calcul : tools/werbomont/engine.py (lecture, géolocalisation) et
         tools/werbomont/itineraires.py (lignes, PTV, passage par Werbomont).
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import sys
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import ui  # noqa: E402
from tools.werbomont import engine as E  # noqa: E402
from tools.werbomont import itineraires as I  # noqa: E402

ui.page("werbomont", "Les itinéraires chargement → déchargement de chaque ligne, et s’ils passent par Werbomont.")

st.markdown("""
<style>
.wb-kpis { display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; margin: .4rem 0 1rem; }
.wb-kpi { background: #fff; border-radius: 18px; padding: 16px 18px; box-shadow: var(--ap-shadow); }
.wb-kpi .v { font-size: 30px; font-weight: 600; letter-spacing: -0.02em; line-height: 1.1; }
.wb-kpi .l { font-size: 13px; color: var(--ap-sub); margin-top: 4px; }
.wb-kpi .s { font-size: 12px; color: var(--ap-sub); margin-top: 2px; }
@media (max-width: 760px) { .wb-kpis { grid-template-columns: repeat(2, 1fr); } }
</style>
""", unsafe_allow_html=True)


# ─── Utilitaires ─────────────────────────────────────────────────────────────
def fr2(x) -> str:
    return f"{x:.2f}".replace(".", ",")


def nb(x) -> str:
    return f"{x:_.0f}".replace("_", " ")


def empreinte(b: bytes) -> str:
    return hashlib.md5(b).hexdigest()


@st.cache_resource(show_spinner=False)
def geocodeur():
    return E.Geocodeur()


@st.cache_resource(show_spinner=False)
def cache_itineraires() -> dict:
    """Itinéraires PTV partagés entre sessions tant que l'app tourne."""
    return {}


@st.cache_data(show_spinner=False, max_entries=3)
def lire_excel(contenu: bytes) -> pd.DataFrame:
    return pd.read_excel(io.BytesIO(contenu), dtype={"Code postal": str, "Département": str})


@st.cache_data(show_spinner=False, max_entries=3)
def preparer(h_mis: str, h_ca: str, _mis: bytes, _ca: bytes):
    m = E.preparer_missions(lire_excel(_mis), geocodeur())
    ca = E.preparer_ca(lire_excel(_ca))
    return E.construire_dossiers(m, ca)


BORDER_URLS = [
    "https://raw.githubusercontent.com/nvkelso/natural-earth-vector/master/geojson/ne_10m_admin_0_boundary_lines_land.geojson",
    "https://raw.githubusercontent.com/nvkelso/natural-earth-vector/master/geojson/ne_50m_admin_0_boundary_lines_land.geojson",
]


@st.cache_data(ttl=30 * 86400, show_spinner="Chargement des frontières…")
def load_borders():
    """Frontières Natural Earth limitées à l'Europe (même source que la carte des lavages)."""
    import urllib.request as ureq
    w, s_, e, n = (-11.0, 35.0, 32.0, 62.0)
    for url in BORDER_URLS:
        try:
            req = ureq.Request(url, headers={"User-Agent": "CB-Transport-Hub/1.0"})
            with ureq.urlopen(req, timeout=40) as r:
                data = json.loads(r.read())
        except Exception:
            continue
        lignes = []
        for f in data.get("features", []):
            g = f.get("geometry") or {}
            parts = [g["coordinates"]] if g.get("type") == "LineString" else (
                g["coordinates"] if g.get("type") == "MultiLineString" else [])
            for li in parts:
                if any(w <= p[0] <= e and s_ <= p[1] <= n for p in li):
                    lignes.append(I._decimer([[p[1], p[0]] for p in li], 0.003))
        if lignes:
            lignes = [[[p[1], p[0]] for p in li] for li in lignes]      # retour en [lon, lat] (GeoJSON)
            return {"type": "FeatureCollection", "features": [{"type": "Feature", "properties": {},
                    "geometry": {"type": "MultiLineString", "coordinates": lignes}}]}
    raise RuntimeError("frontières indisponibles")


# ─── Fichiers et paramètres ──────────────────────────────────────────────────
ui.section("Fichiers")
c1, c2 = st.columns(2)
with c1:
    f_mis = st.file_uploader("Liste des missions (LISTES_MISSIONS)", type=["xlsx", "xls"], key="wb_mis",
                             help="Une ligne par activité : CHARGER, DECHARGER, LAVAGE…")
with c2:
    f_ca = st.file_uploader("Chiffre d’affaires citernes (CA CIT)", type=["xlsx", "xls"], key="wb_ca",
                            help="Une ligne par dossier avec le prix transport.")

with st.expander("Paramètres"):
    p1, p2, p3 = st.columns(3)
    with p1:
        w_lat = st.number_input("Werbomont — latitude", value=E.WERBOMONT[0], format="%.4f", step=0.001)
        w_lon = st.number_input("Werbomont — longitude", value=E.WERBOMONT[1], format="%.4f", step=0.001)
    with p2:
        s_ok = st.number_input("Détour acceptable jusqu’à (km)", value=E.SEUILS[0], min_value=0, step=5)
        s_verif = st.number_input("À vérifier jusqu’à (km)", value=E.SEUILS[1], min_value=0, step=5)
    with p3:
        couloir = st.number_input("« Passe par Werbomont » si la route passe à moins de (km)",
                                  value=I.COULOIR_KM, min_value=0.5, max_value=30.0, step=0.5)
        exclure_aff = st.toggle("Exclure les affrétés (tracteur AFF)", value=True)

if not (f_mis and f_ca):
    st.markdown("""
<div class="ap-card"><h4>Ce que fait l’outil</h4>
<ul class="ap-list">
<li><b>Lignes</b> <span>— les dossiers sont regroupés par itinéraire chargé : chargement(s) → déchargement(s),
dossiers multi-points compris. Le prix vient du CA CIT, joint par n° de dossier.</span></li>
<li><b>Itinéraire PTV</b> <span>— pour chaque ligne, la route poids lourd réelle est calculée et tracée sur la carte.</span></li>
<li><b>Werbomont</b> <span>— si la route passe à moins de 5 km : elle passe par Werbomont. Sinon, PTV recalcule
la route en passant par Werbomont et donne le détour : ≤ 60 km acceptable, 60–80 à vérifier, au-delà trop lourd.</span></li>
<li><b>Rentabilité</b> <span>— Σ prix transport ÷ Σ km de la ligne, en direct puis via Werbomont.</span></li>
</ul></div>""", unsafe_allow_html=True)
    st.stop()

b_mis, b_ca = f_mis.getvalue(), f_ca.getvalue()
W = (float(w_lat), float(w_lon))
seuils = (float(s_ok), float(max(s_verif, s_ok)))

try:
    with st.spinner("Lecture et géolocalisation des dossiers… (≈ 30 s la première fois)"):
        dossiers = preparer(empreinte(b_mis), empreinte(b_ca), b_mis, b_ca)
except ValueError as e:
    st.error(str(e))
    st.stop()

dmin_, dmax_ = dossiers["dt_C1"].min(), dossiers["dt_C1"].max()
f1, f2 = st.columns([2, 1])
with f1:
    periode = st.date_input("Période (date de chargement)", value=(dmin_.date(), dmax_.date()),
                            min_value=dmin_.date(), max_value=dmax_.date(), format="DD/MM/YYYY")
with f2:
    min_dos = st.number_input("Lignes d’au moins … dossiers", min_value=1, max_value=200, value=5,
                              help="Les lignes rares sont écartées pour limiter les appels PTV.")

d = dossiers
if exclure_aff:
    d = d[~d["tracteur"].fillna("").str.upper().str.startswith("AFF")]
if isinstance(periode, tuple) and len(periode) == 2:
    d = d[(d["dt_C1"] >= pd.Timestamp(periode[0])) & (d["dt_C1"] < pd.Timestamp(periode[1]) + pd.Timedelta(days=1))]
lignes = I.construire_lignes(d, int(min_dos))
cache = cache_itineraires()

# ─── Itinéraires PTV ─────────────────────────────────────────────────────────
ui.section("Itinéraires PTV")
cle = E.cle_ptv()
reste = I.a_calculer(lignes, cache, W, seuils, couloir)
faits = len(lignes) - reste
st.markdown(
    (f'<div class="ap-statut"><i></i>PTV configuré · {faits} lignes sur {len(lignes)} calculées</div>' if cle else
     '<div class="ap-statut off"><i></i>Clé PTV non configurée (PTV_API_KEY) : aucun itinéraire possible</div>'),
    unsafe_allow_html=True)
st.caption(f"{len(lignes)} lignes de {int(min_dos)} dossiers ou plus, soit {nb(lignes['nb_dossiers'].sum())} dossiers "
           f"sur {nb(len(d))}. Un appel PTV par ligne, plus un second pour celles proches de Werbomont.")

if cle:
    cT, cG = st.columns([1, 2])
    with cT:
        if st.button("Tester PTV"):
            with st.spinner("Liège → Werbomont…"):
                km_t, err_t = I.ptv_test(cle)
            (st.success(f"PTV répond : Liège → Werbomont = {fr2(km_t)} km.") if km_t
             else st.error(f"PTV ne répond pas correctement : {err_t}"))
    if reste:
        a1, a2, a3 = st.columns([1, 1, 2])
        with a1:
            lot = st.number_input("Lignes par lancement", min_value=5, max_value=2000, value=min(100, reste), step=25)
        with a2:
            paral = st.number_input("Appels simultanés", min_value=1, max_value=6, value=2,
                                    help="Baissez à 1 si PTV renvoie des erreurs 429 (trop de requêtes).")
        with a3:
            st.write("")
            st.write("")
            go = st.button(f"Calculer {min(int(lot), reste)} itinéraires", type="primary")
        if go:
            barre = st.progress(0.0, text="PTV…")
            ok, erreurs = I.calculer_lot(
                lignes, cle, cache, W, seuils, couloir, limite=int(lot), workers=int(paral),
                progression=lambda n, t, o, e: barre.progress(n / t, text=f"PTV… {n}/{t} · {o} ok · {e} en échec"))
            barre.empty()
            st.session_state["wb_bilan"] = (ok, erreurs)
            st.rerun()
    if st.session_state.get("wb_bilan"):
        ok_b, err_b = st.session_state["wb_bilan"]
        if err_b:
            detail = " · ".join(f"{k} ({v})" for k, v in sorted(err_b.items(), key=lambda x: -x[1]))
            (st.warning if ok_b else st.error)(f"Dernier lot : {ok_b} lignes calculées, {sum(err_b.values())} en échec. "
                                               f"Erreurs PTV : {detail}")
        else:
            st.success(f"Dernier lot : {ok_b} lignes calculées.")

with st.expander("Sauvegarder ou recharger les itinéraires déjà calculés"):
    st.caption("Streamlit Cloud oublie les itinéraires quand l’app se met en veille. Téléchargez-les après un calcul, "
               "rechargez-les la fois suivante : aucun appel PTV à refaire.")
    s1, s2 = st.columns(2)
    with s1:
        st.download_button("Télécharger les itinéraires (JSON)", json.dumps(cache).encode("utf-8"),
                           f"itineraires_werbomont_{datetime.now():%Y%m%d}.json", "application/json",
                           disabled=not cache)
    with s2:
        f_cache = st.file_uploader("Recharger un fichier d’itinéraires", type=["json"], key="wb_cache")
        if f_cache is not None and st.session_state.get("wb_cache_lu") != f_cache.file_id:
            try:
                cache.update(json.loads(f_cache.getvalue()))
                st.session_state["wb_cache_lu"] = f_cache.file_id
                st.rerun()
            except Exception as e:
                st.error(f"Fichier illisible : {e}")

# ─── Résultats ───────────────────────────────────────────────────────────────
R = I.resultats(lignes, cache, W, seuils, couloir)

ui.section("Synthèse")
cnt_l = R["categorie"].value_counts()
cnt_d = R.groupby("categorie")["nb_dossiers"].sum()
tot_d = R["nb_dossiers"].sum() or 1


def kpi(k, label):
    lib, coul = I.CAT[k]
    return (f'<div class="wb-kpi"><div class="v" style="color:{coul}">{int(cnt_l.get(k, 0))}</div>'
            f'<div class="l">{label}</div><div class="s">{nb(cnt_d.get(k, 0))} dossiers · '
            f'{cnt_d.get(k, 0) / tot_d * 100:.0f} %</div></div>')


st.markdown('<div class="wb-kpis">' + kpi("passe", "lignes passent par Werbomont")
            + kpi("ok", f"détour ≤ {seuils[0]:.0f} km") + kpi("verif", f"détour {seuils[0]:.0f}–{seuils[1]:.0f} km")
            + kpi("lourd", f"détour > {seuils[1]:.0f} km") + "</div>", unsafe_allow_html=True)
if cnt_l.get("na", 0):
    st.markdown(f'<div class="ap-note">{int(cnt_l["na"])} lignes ({nb(cnt_d.get("na", 0))} dossiers) n’ont pas encore '
                f'd’itinéraire PTV : elles apparaissent en pointillés gris sur la carte.</div>', unsafe_allow_html=True)
est = R[(R["source_detour"] == "estimé")]
if len(est):
    st.caption(f"{len(est)} lignes passent à plus de {seuils[1] + I.MARGE_ESTIMATION:.0f} km de détour estimé : "
               "classées trop lourdes sans second appel PTV (détour estimé au vol d’oiseau × 1,30).")


# ─── Carte ───────────────────────────────────────────────────────────────────
def _v(x, nd=2):
    if x is None:
        return None
    try:
        if pd.isna(x):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(x, (np.integer, int)):
        return int(x)
    if isinstance(x, (np.floating, float)):
        return round(float(x), nd)
    return x


def payload_carte() -> dict:
    rows = []
    for r in R.itertuples(index=False):
        e = cache.get(r.cle) or {}
        rows.append({
            "l": r.ligne, "cli": r.client, "n": int(r.nb_dossiers), "prix": _v(r.prix_moyen, 0), "ca": _v(r.ca_total, 0),
            "km": _v(r.km_direct, 0), "kmw": _v(r.km_via_w, 0), "det": _v(r.detour_km, 0), "src": r.source_detour,
            "dw": _v(r.dist_w_km, 1), "cat": r.categorie, "rd": _v(r.renta_directe), "rw": _v(r.renta_via_w),
            "dp": _v(r.ecart_renta_pct, 1), "ins": r.insertion, "deja": int(r.deja_relais_W),
            "pts": [[round(p[0], 4), round(p[1], 4)] for p in r.points],
            "tr": e.get("trace"), "trw": e.get("trace_w"),
        })
    per = (f"{pd.Timestamp(periode[0]):%d/%m/%Y} – {pd.Timestamp(periode[1]):%d/%m/%Y}"
           if isinstance(periode, tuple) and len(periode) == 2 else "")
    return {"lignes": rows, "w": list(W), "seuils": list(seuils), "couloir": couloir, "periode": per,
            "cat": {k: list(v) for k, v in I.CAT.items()}, "genere": datetime.now().strftime("%d/%m/%Y à %H:%M")}


ui.section("Carte")
try:
    frontieres = load_borders()
except Exception:
    frontieres = None
MAP_HTML = open(os.path.join(os.path.dirname(I.__file__), "carte.html"), encoding="utf-8").read()
carte = (MAP_HTML.replace("__DATA__", json.dumps(payload_carte(), ensure_ascii=False).replace("</", "<\\/"), 1)
         .replace("__BORDERS__", json.dumps(frontieres, separators=(",", ":")) if frontieres else "null", 1))
contenu = json.dumps(carte).replace("</", "<\\/")
components.html(f"""
<style>
  body {{ margin: 0; font-family: -apple-system, BlinkMacSystemFont, "SF Pro Text", "Inter", "Helvetica Neue", Arial, sans-serif; }}
  button {{ height: 48px; padding: 0 28px; border: none; border-radius: 980px; background: #0071e3; color: #fff;
           font: 500 17px -apple-system, BlinkMacSystemFont, "SF Pro Text", "Inter", "Helvetica Neue", Arial, sans-serif;
           cursor: pointer; transition: background .15s, transform .1s; }}
  button:hover {{ background: #0077ed; }} button:active {{ transform: scale(.98); }}
  p {{ margin: 10px 0 0 4px; font-size: 13px; color: #6e6e73; }}
</style>
<button id="ouvrir">Ouvrir la carte plein écran</button>
<p id="msg">S’ouvre dans un nouvel onglet.</p>
<script>
const PAGE = {contenu};
document.getElementById('ouvrir').addEventListener('click', () => {{
  const url = URL.createObjectURL(new Blob([PAGE], {{ type: 'text/html' }}));
  const w = window.open(url, '_blank');
  document.getElementById('msg').textContent = w ? 'Carte ouverte dans un nouvel onglet.'
    : 'Le navigateur a bloqué l’ouverture : autorisez les pop-up pour le HUB, ou téléchargez la carte ci-dessous.';
}});
</script>""", height=90)
st.download_button("Télécharger la carte (HTML)", carte.encode("utf-8"), "carte_werbomont.html", "text/html")

# ─── Tableau ─────────────────────────────────────────────────────────────────
ui.section("Lignes")
LIB = {k: v[0] for k, v in I.CAT.items()}
cats = st.multiselect("Catégories", list(LIB.values()), default=[LIB["passe"], LIB["ok"], LIB["verif"]], key="wb_cats")
T = R[R["categorie"].map(LIB).isin(cats)].copy()
T["categorie"] = T["categorie"].map(LIB)
COLS = ["ligne", "client", "nb_dossiers", "prix_moyen", "km_direct", "dist_w_km", "detour_km", "source_detour",
        "km_via_w", "insertion", "renta_directe", "renta_via_w", "ecart_renta_pct", "deja_relais_W", "categorie"]
st.dataframe(T[COLS], hide_index=True, width="stretch", height=480, column_config={
    "ligne": st.column_config.TextColumn("Ligne", width="large"), "client": "Client",
    "nb_dossiers": st.column_config.NumberColumn("Dossiers"),
    "prix_moyen": st.column_config.NumberColumn("Prix moy.", format="%.0f €"),
    "km_direct": st.column_config.NumberColumn("Km PTV", format="%.0f"),
    "dist_w_km": st.column_config.NumberColumn("Route ↔ W", format="%.1f km"),
    "detour_km": st.column_config.NumberColumn("Détour", format="%.0f km"),
    "source_detour": "Source détour",
    "km_via_w": st.column_config.NumberColumn("Km via W", format="%.0f"),
    "insertion": "W inséré entre",
    "renta_directe": st.column_config.NumberColumn("Renta", format="%.2f €/km"),
    "renta_via_w": st.column_config.NumberColumn("Renta via W", format="%.2f €/km"),
    "ecart_renta_pct": st.column_config.NumberColumn("Écart", format="%.1f %%"),
    "deja_relais_W": st.column_config.NumberColumn("Relais W déjà faits"),
    "categorie": "Catégorie"})


# ─── Export Excel ────────────────────────────────────────────────────────────
def export_excel() -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fonds = {"Passe par Werbomont": "E5F1FF", "Détour acceptable": "E8F6EC", "Détour à vérifier": "FFF1E0",
             "Détour trop lourd": "FDECEE", "Pas encore calculé": "F2F2F7"}
    X = R.copy()
    X["categorie"] = X["categorie"].map(LIB)
    X = X[COLS].rename(columns={
        "ligne": "Ligne", "client": "Client", "nb_dossiers": "Dossiers", "prix_moyen": "Prix moyen €",
        "km_direct": "Km PTV", "dist_w_km": "Route ↔ Werbomont km", "detour_km": "Détour km",
        "source_detour": "Source détour", "km_via_w": "Km via Werbomont", "insertion": "Werbomont inséré entre",
        "renta_directe": "Renta €/km", "renta_via_w": "Renta via W €/km", "ecart_renta_pct": "Écart renta %",
        "deja_relais_W": "Relais W déjà faits", "categorie": "Catégorie"})
    synth = pd.DataFrame([
        ("Période", payload_carte()["periode"]), ("Werbomont (lat, lon)", f"{W[0]:.4f}, {W[1]:.4f}"),
        ("Passe par Werbomont", f"itinéraire PTV à moins de {couloir:g} km"),
        ("Seuils détour", f"acceptable ≤ {seuils[0]:.0f} km · à vérifier ≤ {seuils[1]:.0f} km"),
        ("Lignes", f"{len(R)} lignes d'au moins {int(min_dos)} dossiers"),
        ("Rentabilité", "Σ prix transport ÷ Σ km chargés de la ligne (PTV), direct puis via Werbomont"),
        ("Généré le", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ], columns=["Élément", "Valeur"])
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        synth.to_excel(xw, sheet_name="Synthèse", index=False)
        X.to_excel(xw, sheet_name="Lignes", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(color="FFFFFF", bold=True)
                c.fill = PatternFill("solid", fgColor="003087")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "B2"
            ws.row_dimensions[1].height = 30
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), 1):
                larg = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, larg + 2), 60)
        ws = xw.book["Lignes"]
        entetes = [c.value for c in ws[1]]
        for j, h in enumerate(entetes, 1):
            fmt = "0.00" if "Renta" in h else ("0.0" if "%" in h else ("#,##0" if ("km" in h.lower() or "€" in h) else None))
            for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                if fmt:
                    cell.number_format = fmt
                if h == "Catégorie" and cell.value in fonds:
                    cell.fill = PatternFill("solid", fgColor=fonds[cell.value])
        ws.auto_filter.ref = ws.dimensions
        xw.book["Synthèse"].column_dimensions["B"].width = 90
    return out.getvalue()


st.download_button("Télécharger les lignes (Excel)", export_excel(),
                   f"itineraires_werbomont_{datetime.now():%Y%m%d}.xlsx",
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
