"""
Page Streamlit : Optimiseur de groupage (version test)
Associe les commandes à faire en camions groupés : capacité, dates, nombre d'arrêts et
allongement de chaque lot respectés, ordre des arrêts optimisé.

Fichiers (pour tester : l'export d'une période) : LISTES_MISSIONS et/ou export CA,
joints par n° de dossier. Calcul : tools/groupage/optimiseur.py.
"""

from __future__ import annotations

import hashlib
import io
import os
import sys
from datetime import datetime

import numpy as np
import pandas as pd
import pydeck as pdk
import streamlit as st

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import ui  # noqa: E402
from tools.crossdock import engine as X  # noqa: E402
from tools.groupage import optimiseur as O  # noqa: E402
from tools.werbomont import engine as W  # noqa: E402

ui.page("groupage", "Les commandes à faire, associées en camions groupés : moins de camions, moins de km.")

st.markdown("""
<style>
.og-kpis { display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; margin: .4rem 0 1rem; }
.og-kpi { background: #fff; border-radius: 18px; padding: 16px 18px; box-shadow: var(--ap-shadow); }
.og-kpi .v { font-size: 30px; font-weight: 600; letter-spacing: -0.02em; line-height: 1.1; color: var(--ap-text); }
.og-kpi .l { font-size: 13px; color: var(--ap-sub); margin-top: 4px; }
.og-kpi .s { font-size: 12px; color: var(--ap-sub); margin-top: 2px; }
.og-kpi.pos .v { color: var(--ap-green); }
.og-etapes { list-style: none; margin: 0; padding: 0; counter-reset: e; }
.og-etapes li { counter-increment: e; display: flex; gap: 12px; align-items: baseline; padding: 8px 0;
  border-bottom: .5px solid var(--ap-hair); font-size: 15px; }
.og-etapes li::before { content: counter(e); flex: none; width: 24px; height: 24px; border-radius: 50%;
  display: inline-flex; align-items: center; justify-content: center; font-size: 12px; font-weight: 600; color: #fff; }
.og-etapes li.c::before { background: #0071e3; } .og-etapes li.d::before { background: #ff9500; }
.og-etapes li b { font-weight: 600; } .og-etapes li span { color: var(--ap-sub); }
@media (max-width: 760px) { .og-kpis { grid-template-columns: repeat(2, 1fr); } }
</style>
""", unsafe_allow_html=True)

CARTO_LIGHT = "https://basemaps.cartocdn.com/gl/positron-gl-style/style.json"


def nb(x) -> str:
    return "–" if x is None or pd.isna(x) else f"{x:_.0f}".replace("_", " ")


def eur(x) -> str:
    return "–" if x is None or pd.isna(x) else f"{x:_.0f} €".replace("_", " ")


def kpi(v, l, s="", pos=False) -> str:
    return (f'<div class="og-kpi{" pos" if pos else ""}"><div class="v">{v}</div><div class="l">{l}</div>'
            f'<div class="s">{s}</div></div>')


def vue_carte(pts: pd.DataFrame) -> pdk.ViewState:
    """Cadre la carte sur tous les arrêts (carte de 460 px de haut, ~650 px de large)."""
    la0, la1, lo0, lo1 = pts["lat"].min(), pts["lat"].max(), pts["lon"].min(), pts["lon"].max()
    lat_c = (la0 + la1) / 2
    etendue_lon = max((lo1 - lo0) * 1.3, 0.2)
    etendue_lat = max((la1 - la0) * 1.3, 0.2) / max(np.cos(np.radians(lat_c)), 0.3)
    zoom = min(np.log2(360 * 650 / 512 / etendue_lon), np.log2(360 * 460 / 512 / etendue_lat), 11)
    return pdk.ViewState(latitude=float(lat_c), longitude=float((lo0 + lo1) / 2), zoom=float(zoom))


@st.cache_resource(show_spinner=False)
def geocodeur():
    return W.Geocodeur()


@st.cache_data(show_spinner=False, max_entries=3)
def lire_excel(contenu: bytes) -> pd.DataFrame:
    return pd.read_excel(io.BytesIO(contenu), dtype={"Code postal": str, "C.P. chargement": str,
                                                     "C.P. déchargement": str, "N° Dossier": str})


@st.cache_data(show_spinner=False, max_entries=3)
def preparer(h_mis: str, h_ca: str, _mis: bytes | None, _ca: bytes | None):
    raw_m = lire_excel(_mis) if _mis else None
    raw_c = lire_excel(_ca) if _ca else None
    if raw_m is not None:
        m = X.preparer_missions(raw_m, geocodeur())
        ca = X.preparer_ca(raw_c) if raw_c is not None else None
        cmds = O.commandes_depuis_missions(m, ca)
    else:
        cmds = O.commandes_depuis_ca(raw_c, geocodeur())
    qte, source = (None, "")
    for raw in (raw_m, raw_c):
        if raw is not None and qte is None:
            qte, source = O.quantite_pct(raw)
    return cmds, qte, source


@st.cache_data(show_spinner=False, max_entries=6)
def optimiser(cle: str, _cmds: pd.DataFrame, params: dict):
    return O.optimiser(_cmds, params)


# ─── Fichiers ────────────────────────────────────────────────────────────────
ui.section("Commandes")
c1, c2 = st.columns(2)
with c1:
    f_mis = st.file_uploader("Liste des missions (LISTES_MISSIONS)", type=["xlsx", "xls"], key="og_mis",
                             help="Arrêts réels de chaque dossier (multi-points compris).")
with c2:
    f_ca = st.file_uploader("Chiffre d’affaires (export CA)", type=["xlsx", "xls"], key="og_ca",
                            help="Seul, il suffit : une adresse de chargement et de déchargement par dossier. "
                                 "Avec les missions, il apporte le prix et le client.")

if not (f_mis or f_ca):
    st.markdown("""
<div class="ap-card"><h4>Ce que fait l’outil</h4>
<ul class="ap-list">
<li><b>Commandes</b> <span>— chaque dossier de la période choisie devient une commande à placer : ses
chargements, ses déchargements, sa date et sa taille en part de remorque.</span></li>
<li><b>Associations</b> <span>— les commandes sont regroupées dans un même camion quand ça fait gagner des
km, sans dépasser la remorque, l’écart de dates, le nombre d’arrêts ni l’allongement admis pour chaque lot.</span></li>
<li><b>Tournée</b> <span>— pour chaque camion : l’ordre des chargements puis des livraisons, le plus court trouvé.</span></li>
<li><b>Résultat</b> <span>— camions et km avant/après, plan camion par camion, carte et export Excel.</span></li>
</ul></div>
<div class="ap-note">Version test : on rejoue une période passée comme si c’étaient les commandes à faire.
Plus tard, la même page pourra lire les commandes ouvertes de TAS. Km au vol d’oiseau × 1,30, sans PTV.</div>""",
                unsafe_allow_html=True)
    st.stop()

b_mis = f_mis.getvalue() if f_mis else None
b_ca = f_ca.getvalue() if f_ca else None
try:
    with st.spinner("Lecture et géolocalisation des commandes…"):
        cmds_all, qte, source_qte = preparer(hashlib.md5(b_mis).hexdigest() if b_mis else "",
                                             hashlib.md5(b_ca).hexdigest() if b_ca else "", b_mis, b_ca)
except ValueError as e:
    st.error(str(e))
    st.stop()

ok = cmds_all[cmds_all["statut"].eq("ok") & pd.to_datetime(cmds_all["dt"]).notna()].copy()
if ok.empty:
    st.error("Aucune commande exploitable (adresses non géolocalisées ou dates manquantes).")
    st.stop()
ok["dt"] = pd.to_datetime(ok["dt"])

# ─── Sélection des commandes ─────────────────────────────────────────────────
dmax = ok["dt"].max().normalize()
f1, f2, f3 = st.columns([2, 2, 1])
with f1:
    periode = st.date_input("Commandes à planifier (date de chargement)",
                            value=((dmax - pd.Timedelta(days=6)).date(), dmax.date()),
                            min_value=ok["dt"].min().date(), max_value=dmax.date(), format="DD/MM/YYYY",
                            help="Pour tester : une semaine de l’export, rejouée comme des commandes à faire.")
with f2:
    types_tr = sorted(t for t in ok["type_transport"].fillna("").astype(str).unique() if t)
    sel_types = st.multiselect("Type de transport", types_tr, default=[], placeholder="Tous")
with f3:
    excl_aff = st.toggle("Exclure les affrétés", value=True)

c = ok
if isinstance(periode, tuple) and len(periode) == 2:
    c = c[(c["dt"] >= pd.Timestamp(periode[0])) & (c["dt"] < pd.Timestamp(periode[1]) + pd.Timedelta(days=1))]
if sel_types:
    c = c[c["type_transport"].isin(sel_types)]
if excl_aff and "tracteur" in c:
    c = c[~c["tracteur"].fillna("").astype(str).str.upper().str.startswith("AFF")]
c = c.copy()
c["client"] = c["client"].fillna("").astype(str).replace("", "(sans client)")
if c.empty:
    st.warning("Aucune commande sur cette période.")
    st.stop()
if len(c) > 1500:
    st.warning(f"{nb(len(c))} commandes : le calcul peut prendre plusieurs minutes. Réduisez la période si besoin.")

# ─── Taille des lots ─────────────────────────────────────────────────────────
ui.section("Taille des lots")
if qte is not None:
    st.markdown(f'<div class="ap-note">Taille lue dans l’export : <b>{source_qte}</b>, convertie en part de '
                f'remorque. Les commandes sans valeur prennent la taille par client ci-dessous.</div>',
                unsafe_allow_html=True)
else:
    st.caption("Les exports ne donnent ni poids ni mètres de plancher : indiquez la part de remorque qu’occupe "
               "en général une commande de chaque client. 100 % = camion complet, jamais regroupé.")
t1, t2 = st.columns([1, 2])
with t1:
    taille_def = st.number_input("Taille par défaut (% de remorque)", 5, 100, 50, 5)
cle_cli = hashlib.md5(("|".join(sorted(c["client"].unique())) + str(taille_def)).encode()).hexdigest()
if st.session_state.get("og_cli_cle") != cle_cli:
    cnt = c.groupby("client").size().sort_values(ascending=False)
    ancien = st.session_state.get("og_tailles", {})
    st.session_state["og_cli_df"] = pd.DataFrame({"client": cnt.index, "commandes": cnt.values,
                                                  "taille": [ancien.get(k, taille_def) for k in cnt.index]})
    st.session_state["og_cli_cle"] = cle_cli
with t2:
    cli_df = st.data_editor(st.session_state["og_cli_df"], hide_index=True, width="stretch", height=240,
                            key=f"og_cli_ed_{cle_cli[:8]}", disabled=["client", "commandes"], column_config={
                                "client": "Client", "commandes": st.column_config.NumberColumn("Commandes"),
                                "taille": st.column_config.NumberColumn("Taille (% remorque)", min_value=1,
                                                                        max_value=100, step=5, format="%d %%")})
tailles = dict(zip(cli_df["client"], cli_df["taille"].fillna(taille_def)))
st.session_state["og_tailles"] = {**st.session_state.get("og_tailles", {}), **tailles}
c["taille"] = c["client"].map(tailles).fillna(taille_def).astype(float)
if qte is not None:
    q = c["dossier"].map(qte)
    c.loc[q.notna(), "taille"] = q[q.notna()].clip(lower=1)

with st.expander(f"Les {nb(len(c))} commandes (taille et inclusion modifiables une par une)"):
    vue = c.assign(route=[f"{(a or [''])[0]} → {(b or [''])[-1]}" for a, b in zip(c["villes_C"], c["villes_D"])],
                   inclure=True)[["inclure", "dossier", "dt", "client", "route", "taille", "prix"]]
    ed = st.data_editor(vue, hide_index=True, width="stretch", height=380, key=f"og_cmd_ed_{cle_cli[:8]}",
                        disabled=["dossier", "dt", "client", "route", "prix"], column_config={
                            "inclure": st.column_config.CheckboxColumn("Inclure"), "dossier": "Dossier",
                            "dt": st.column_config.DatetimeColumn("Chargement", format="DD/MM/YY HH:mm"),
                            "client": "Client", "route": st.column_config.TextColumn("Route", width="large"),
                            "taille": st.column_config.NumberColumn("Taille", min_value=1, max_value=100,
                                                                    format="%.0f %%"),
                            "prix": st.column_config.NumberColumn("Prix", format="%.0f €")})
    c["taille"] = ed["taille"].to_numpy()
    c = c[ed["inclure"].to_numpy()]

# ─── Contraintes ─────────────────────────────────────────────────────────────
with st.expander("Contraintes", expanded=False):
    k1, k2, k3 = st.columns(3)
    with k1:
        fenetre = st.number_input("Écart de dates max. dans un camion (jours)", 0.0, 7.0, 2.0, 0.5,
                                  help="Entre le premier et le dernier chargement d’un même camion.")
        rayon = st.number_input("Chargements à moins de (km)", 10, 400, 100, 10,
                                help="À vol d’oiseau : deux camions ne sont fusionnés que si au moins deux de "
                                     "leurs chargements sont aussi proches.")
    with k2:
        max_lots = st.number_input("Lots max. par camion", 2, 10, 4, 1)
        max_ch = st.number_input("Chargements max.", 1, 10, 4, 1)
        max_de = st.number_input("Déchargements max.", 1, 12, 5, 1)
    with k3:
        allong_pct = st.number_input("Allongement max. d’un lot (%)", 0, 300, 50, 10,
                                     help="Trajet du lot dans le camion groupé (de son chargement à sa livraison) "
                                          "comparé à son trajet direct.")
        allong_km = st.number_input("… plus une marge fixe (km)", 0, 300, 50, 10)
        cout_km = st.number_input("Coût du km (€)", 0.0, 5.0, 1.20, 0.05, format="%.2f")
params = dict(fenetre_jours=float(fenetre), rayon_ch=float(rayon), max_lots=int(max_lots), max_ch=int(max_ch),
              max_de=int(max_de), allong_pct=float(allong_pct), allong_km=float(allong_km), capacite=O.CAPACITE,
              passes=3)

cle = hashlib.md5((",".join(f"{d}:{t:.0f}" for d, t in zip(c["dossier"], c["taille"]))
                   + repr(sorted(params.items()))).encode()).hexdigest()
if st.session_state.get("og_cle") != cle:
    st.session_state.pop("og_res", None)
lancer = st.button(f"Optimiser les groupages ({nb(len(c))} commandes)", type="primary",
                   disabled=len(c) < 2)
if lancer:
    with st.spinner("Recherche des meilleures associations…"):
        st.session_state["og_res"] = optimiser(cle, c.reset_index(drop=True), params)
        st.session_state["og_cle"] = cle
if "og_res" not in st.session_state:
    st.stop()

camions, journal = st.session_state["og_res"]
cc = c.reset_index(drop=True)
C, K = O.plan(cc, camions, float(cout_km))

# ─── Résultat ────────────────────────────────────────────────────────────────
ui.section("Résultat")
grp = C[C["groupage"]]
km_av, km_ap = C["km_seuls"].sum(), C["km"].sum()
st.markdown('<div class="og-kpis">'
            + kpi(f"{nb(len(C))}", "camions au lieu de " + nb(len(cc)), f"{nb(len(cc) - len(C))} camions en moins",
                  pos=len(C) < len(cc))
            + kpi(f"{nb(len(grp))}", "groupages", f"{nb(grp['lots'].sum())} commandes regroupées")
            + kpi(f"−{(1 - km_ap / km_av) * 100:.0f} %" if km_av else "–", "de km chargés",
                  f"{nb(km_av)} → {nb(km_ap)} km", pos=km_ap < km_av)
            + kpi(eur(C["gain_eur"].sum()), "économisés au coût du km",
                  f"remplissage moyen des groupages {grp['remplissage'].mean():.0f} %" if len(grp) else "", pos=True)
            + "</div>", unsafe_allow_html=True)
st.caption("Km chargés seulement (du premier chargement à la dernière livraison) : les trajets à vide ne sont "
           "pas comptés, avant comme après.")

seuls_grp = st.toggle("Seulement les groupages", value=True)
V = grp if seuls_grp else C
st.dataframe(V.drop(columns=["arrets_liste"]), hide_index=True, width="stretch", height=440, column_config={
    "camion": st.column_config.NumberColumn("Camion"), "groupage": None,
    "lots": st.column_config.NumberColumn("Lots"), "dossiers": "Dossiers", "clients": "Clients",
    "date": st.column_config.DatetimeColumn("1er chargement", format="DD/MM HH:mm"),
    "remplissage": st.column_config.ProgressColumn("Remplissage", min_value=0, max_value=100, format="%.0f %%"),
    "chargements": st.column_config.NumberColumn("Charg."), "dechargements": st.column_config.NumberColumn("Livr."),
    "km": st.column_config.NumberColumn("Km", format="%.0f"),
    "km_seuls": st.column_config.NumberColumn("Km séparés", format="%.0f"),
    "km_gagnes": st.column_config.NumberColumn("Gagnés", format="%.0f km"),
    "gain_eur": st.column_config.NumberColumn("Gain", format="%.0f €"),
    "prix": st.column_config.NumberColumn("Prix", format="%.0f €"),
    "eur_km": st.column_config.NumberColumn("€/km", format="%.2f"),
    "itineraire": st.column_config.TextColumn("Tournée", width="large")})

# ─── Détail d’un camion ──────────────────────────────────────────────────────
ui.section("Détail d’un camion")
choix = (grp if len(grp) else C)["camion"].tolist()
lib = {r.camion: f"Camion {r.camion} · {r.lots} lot{'s' if r.lots > 1 else ''} · {r.dossiers}" for r in C.itertuples()}
n_sel = st.selectbox("Camion", choix, format_func=lambda k: lib[k])
R = C[C["camion"] == n_sel].iloc[0]
arr = R["arrets_liste"]
g1, g2 = st.columns([1, 1.4])
with g1:
    items = []
    for typ, dos, _, ville in arr:
        k = cc[cc["dossier"] == dos].iloc[0]
        items.append(f'<li class="{"c" if typ == "C" else "d"}"><div><b>{"Charger" if typ == "C" else "Livrer"} '
                     f'{ville}</b><br><span>{dos} · {k["client"]} · {k["taille"]:.0f} %</span></div></li>')
    st.markdown(f'<ul class="og-etapes">{"".join(items)}</ul>', unsafe_allow_html=True)
    st.caption(f"{nb(R['km'])} km en groupage contre {nb(R['km_seuls'])} km en camions séparés · "
               f"remplissage {R['remplissage']:.0f} %")
with g2:
    pts = pd.DataFrame([{"lat": p[0], "lon": p[1], "n": str(i + 1),
                         "couleur": [0, 113, 227] if t == "C" else [255, 149, 0],
                         "info": f"{i + 1}. {'Charger' if t == 'C' else 'Livrer'} {v} ({d})"}
                        for i, (t, d, p, v) in enumerate(arr)])
    chemin = [{"path": [[p[1], p[0]] for _, _, p, _ in arr]}]
    st.pydeck_chart(pdk.Deck(
        layers=[pdk.Layer("PathLayer", data=chemin, get_path="path", get_color=[29, 29, 31, 170], width_min_pixels=3),
                pdk.Layer("ScatterplotLayer", data=pts, get_position="[lon, lat]", get_fill_color="couleur",
                          get_radius=6000, radius_min_pixels=9, stroked=True, get_line_color=[255, 255, 255],
                          line_width_min_pixels=2, pickable=True),
                pdk.Layer("TextLayer", data=pts, get_position="[lon, lat]", get_text="n", get_size=12,
                          get_color=[255, 255, 255], font_weight=700)],
        map_style=CARTO_LIGHT, tooltip={"text": "{info}"},
        initial_view_state=vue_carte(pts)), height=460)

with st.expander("Toutes les commandes et leur camion"):
    st.dataframe(K, hide_index=True, width="stretch", height=380, column_config={
        "dossier": "Dossier", "camion": "Camion", "groupage": "Groupée",
        "date": st.column_config.DatetimeColumn("Chargement", format="DD/MM HH:mm"), "client": "Client",
        "taille": st.column_config.NumberColumn("Taille", format="%.0f %%"), "route": "Route",
        "avec": "Avec", "prix": st.column_config.NumberColumn("Prix", format="%.0f €")})


# ─── Export ──────────────────────────────────────────────────────────────────
def export_excel() -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    etapes = []
    for r in C.itertuples():
        for i, (t, d, p, v) in enumerate(r.arrets_liste, 1):
            etapes.append({"Camion": r.camion, "Étape": i, "Action": "Charger" if t == "C" else "Livrer",
                           "Lieu": v, "Dossier": d, "Latitude": round(p[0], 5), "Longitude": round(p[1], 5)})
    per = (f"{pd.Timestamp(periode[0]):%d/%m/%Y} – {pd.Timestamp(periode[1]):%d/%m/%Y}"
           if isinstance(periode, tuple) and len(periode) == 2 else "")
    synth = pd.DataFrame([
        ("Commandes", f"{len(cc)} ({per})"), ("Camions", f"{len(C)} dont {len(grp)} groupages"),
        ("Km chargés", f"{km_av:.0f} séparés → {km_ap:.0f} groupés"),
        ("Taille des lots", source_qte or f"par client (défaut {taille_def} %)"),
        ("Contraintes", f"écart ≤ {fenetre} j · chargements à < {rayon} km · ≤ {max_lots} lots, {max_ch} charg., "
                        f"{max_de} livr. · allongement ≤ {allong_pct} % + {allong_km} km"),
        ("Km", "vol d'oiseau × 1,30 (sans PTV), trajets à vide non comptés"),
        ("Généré le", datetime.now().strftime("%d/%m/%Y %H:%M"))], columns=["Élément", "Valeur"])
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        synth.to_excel(xw, sheet_name="Synthèse", index=False)
        C.drop(columns=["arrets_liste"]).to_excel(xw, sheet_name="Camions", index=False)
        pd.DataFrame(etapes).to_excel(xw, sheet_name="Tournées", index=False)
        K.to_excel(xw, sheet_name="Commandes", index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="003087")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "B2"
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
                larg = max(len(str(x.value)) if x.value is not None else 0 for x in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, larg + 2), 70)
            for row in ws.iter_rows(min_row=2):
                for x in row:
                    if isinstance(x.value, float) and ws.title != "Tournées":
                        x.number_format = "#,##0"
            ws.auto_filter.ref = ws.dimensions
        xw.book["Synthèse"].column_dimensions["B"].width = 100
    return out.getvalue()


st.download_button("Télécharger le plan de groupage (Excel)", export_excel(),
                   f"plan_groupage_{datetime.now():%Y%m%d}.xlsx",
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
