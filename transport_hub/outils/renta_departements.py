"""
renta_departements.py
──────────────────────────────────────────────────────────────────
Rentabilité et km à vide par département (France) — plateau
──────────────────────────────────────────────────────────────────
Même calcul que « Rentabilité plateau » (exports Traject + CA, cargaison à
bord, km PTV par trajet), puis ventilation par zone de déchargement :
département pour la France, pays ailleurs. Moteur : tools/renta_dept/engine.py
(définitions détaillées en tête de ce fichier).

Le calcul de la page Rentabilité plateau peut être repris tel quel s'il a été
fait dans la session : pas de second passage PTV.
──────────────────────────────────────────────────────────────────
"""

import io
import os
import sys

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import ui  # noqa: E402
from tools.renta_dept import engine as E  # noqa: E402
from tools.renta_dept.plateau import moteur  # noqa: E402

P = moteur()

ui.page("renta_departements", "Renta et km à vide par département de livraison, sur le calcul de "
                              "Rentabilité plateau : vide de sortie, vide d'entrée, rechargement local.")

if st.session_state.get("rd_schema") != E.SCHEMA:
    for k in [k for k in st.session_state if k.startswith("rd_")]:
        st.session_state.pop(k, None)
    st.session_state["rd_schema"] = E.SCHEMA

VERT, ROUGE = "#34a853", "#d93025"

st.markdown("""
<style>
.rd-kpis { display: grid; grid-template-columns: repeat(5, 1fr); gap: 12px; margin: .4rem 0 1rem; }
.rd-kpi { background: #fff; border-radius: 18px; padding: 16px 18px; box-shadow: var(--ap-shadow); }
.rd-kpi .v { font-size: 26px; font-weight: 600; letter-spacing: -0.02em; line-height: 1.1; color: var(--ap-text);
             white-space: nowrap; }
.rd-kpi .l { font-size: 13px; color: var(--ap-sub); margin-top: 4px; }
.rd-kpi .s { font-size: 12px; color: var(--ap-sub); margin-top: 2px; }
@media (max-width: 1100px) { .rd-kpis { grid-template-columns: repeat(3, 1fr); } }
@media (max-width: 700px) { .rd-kpis { grid-template-columns: repeat(2, 1fr); } }
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════
#  AFFICHAGE
# ══════════════════════════════════════════════════════════════════

COLS_ZONES = {
    "libelle": "Département / pays", "livraisons": "Livraisons", "clients": "Clients",
    "pct_groupage": "% groupage", "ca": "CA (€)", "pct_ca": "% du CA", "km_charge": "KM chargés",
    "km_ch_moy": "KM chargés / livraison",
    "renta_chargee": "Renta chargée", "renta_conv": "Renta (convention)", "renta_sortie": "Renta livraison + sortie",
    "sorties": "Repositionnements sortants", "vide_sortie_moy": "Vide de sortie moyen (km)",
    "km_sortie_med": "Vide de sortie médian (km)", "pct_local": "% rechargement local",
    "pct_meme_zone": "% rechargé dans le dépt", "delai_h": "Délai avant rechargement (h)",
    "chargements": "Chargements", "entrees": "Repositionnements entrants",
    "vide_entree_moy": "Vide d'entrée moyen (km)",
    "livraisons_aff": "Livraisons affrétés", "ca_aff": "CA affrétés (€)", "representatif": "Représentatif",
}
COLS_REP = {
    "tracteur": "Tracteur", "statut": "Statut", "dossier_prec": "Dossier livré", "de_ville": "Départ à vide",
    "zone_de_l": "Zone de départ", "dossier_suiv": "Dossier suivant", "a_ville": "Rechargement",
    "zone_a_l": "Zone de rechargement", "km": "KM à vide", "nb_trajets": "Trajets", "delai_h": "Délai (h)",
    "de_ts": "Départ", "a_ts": "Arrivée", "affrete": "Affrété",
}
COLS_CTRL = {"type": "Contrôle", "zone": "Zone", "tracteur": "Tracteur", "dossier": "N° dossier", "detail": "Détail"}

FMT = {
    "livraisons": "%.1f", "pct_groupage": "%.0f %%", "ca": "%.0f €", "pct_ca": "%.1f %%", "km_charge": "%.0f",
    "km_ch_moy": "%.0f", "renta_chargee": "%.2f", "renta_conv": "%.2f", "renta_sortie": "%.2f",
    "vide_sortie_moy": "%.0f", "km_sortie_med": "%.0f", "pct_local": "%.0f %%", "pct_meme_zone": "%.0f %%",
    "delai_h": "%.1f", "vide_entree_moy": "%.0f", "livraisons_aff": "%.1f", "ca_aff": "%.0f €",
}


def _col_config(cols):
    cfg = {}
    for c in cols:
        lab = COLS_ZONES.get(c, COLS_REP.get(c, c))
        if c in FMT:
            cfg[c] = st.column_config.NumberColumn(lab, format=FMT[c])
        elif c in ("de_ts", "a_ts"):
            cfg[c] = st.column_config.DatetimeColumn(lab, format="DD/MM/YY HH:mm")
        elif c in ("representatif", "affrete"):
            cfg[c] = st.column_config.CheckboxColumn(lab)
        else:
            cfg[c] = st.column_config.Column(lab)
    return cfg


def _teinte(v, ref, plus_haut_mieux=True, ecart=0.25):
    """Fond rouge → vert selon l'écart à la moyenne flotte (± ecart relatif = couleur pleine)."""
    if pd.isna(v) or ref is None or pd.isna(ref) or not ref:
        return ""
    x = (v - ref) / abs(ref) / ecart
    x = max(-1.0, min(1.0, x if plus_haut_mieux else -x))
    coul = VERT if x > 0 else ROUGE
    a = int(abs(x) * 0.32 * 255)
    return f"background-color: {coul}{a:02x}" if a > 6 else ""


def styler(df, tot):
    """Renta et vide comparés à la moyenne France (ratio des sommes)."""
    refs = {
        "renta_chargee": (tot.get("renta_chargee"), True),
        "renta_conv": (tot.get("renta_conv"), True),
        "renta_sortie": (tot.get("renta_conv"), True),
        "vide_sortie_moy": (tot.get("vide_sortie_moy"), False),
        "vide_entree_moy": (tot.get("vide_sortie_moy"), False),
        "pct_local": (tot.get("pct_local"), True),
    }
    s = df.style
    for c, (ref, mieux) in refs.items():
        if c in df.columns:
            s = s.map(lambda v, r=ref, m=mieux: _teinte(v, r, m), subset=[c])
    return s


def _ville(l):
    return f"{l[0]} ({l[2]}-{l[1]})" if isinstance(l, (tuple, list)) and len(l) >= 3 else ""


def rep_affichage(rep):
    d = rep.copy()
    d["de_ville"] = d["de"].map(_ville)
    d["a_ville"] = d["a"].map(_ville)
    d["zone_de_l"] = d["zone_de"].map(E.zone_label)
    d["zone_a_l"] = d["zone_a"].map(E.zone_label)
    return d


# ══════════════════════════════════════════════════════════════════
#  EXPORT EXCEL
# ══════════════════════════════════════════════════════════════════

def _style(ws, n):
    head = PatternFill("solid", fgColor="003087")
    alt = PatternFill("solid", fgColor="EEF2F7")
    for cell in ws[1]:
        cell.fill, cell.font = head, Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(2, n + 2):
        if r % 2 == 0:
            for cell in ws[r]:
                cell.fill = alt
    for col in ws.columns:
        L = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(L + 3, 9), 50)
    ws.freeze_panes = "B2"


def export_excel(zones, rep, ctrl, meta):
    out = io.BytesIO()
    fr, hors = zones[zones["france"]], zones[~zones["france"]]
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        pd.DataFrame({"Paramètre": list(meta), "Valeur": list(meta.values())}).to_excel(
            w, sheet_name="Paramètres", index=False)
        _style(w.sheets["Paramètres"], len(meta))
        for nom, df, cols in [("Départements", fr, COLS_ZONES), ("Hors France", hors, COLS_ZONES),
                              ("Repositionnements", rep_affichage(rep), COLS_REP), ("Contrôles", ctrl, COLS_CTRL)]:
            if df is None or df.empty:
                continue
            c = [k for k in cols if k in df.columns]
            x = df[c].rename(columns=cols).copy()
            for col in x.columns:
                if pd.api.types.is_datetime64_any_dtype(x[col]):
                    x[col] = x[col].dt.strftime("%d/%m/%Y %H:%M").fillna("")
                if x[col].dtype == bool:
                    x[col] = x[col].map({True: "Oui", False: "Non"})
            x.to_excel(w, sheet_name=nom, index=False)
            _style(w.sheets[nom], len(x))
    return out.getvalue()


# ══════════════════════════════════════════════════════════════════
#  SOURCE DU CALCUL
# ══════════════════════════════════════════════════════════════════

ui.section("Données")
rp = st.session_state.get("rp_res")
choix = ["Nouveaux fichiers"]
if rp:
    s = rp["sel"]["ts_plan"]
    choix.insert(0, f"Reprendre le calcul de Rentabilité plateau "
                    f"({s.min():%d/%m/%Y} → {s.max():%d/%m/%Y}, {rp['sel']['tracteur'].nunique()} ressources)")
source = st.radio("Source", choix, horizontal=True, label_visibility="collapsed") if len(choix) > 1 else choix[0]

if source.startswith("Reprendre"):
    sel = rp["sel"]
    df_traj, dt, _flotte, ctrl_pl = rp["res"]
    convention = rp.get("convention", "precedent")
    calcul = dict(sel=sel, df_traj=df_traj, dt=dt, ctrl=ctrl_pl, convention=convention,
                  genere=rp.get("genere", ""), source="Rentabilité plateau")
else:
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("#### Exports planning")
        st.caption("Fichiers Traject_xxxxx.xlsx, plusieurs mois si possible : en dessous de 5 livraisons, "
                   "un département n'est pas représentatif.")
        files_pl = st.file_uploader("Plannings (.xlsx)", type=["xlsx"], accept_multiple_files=True,
                                    key="rd_up_pl", label_visibility="collapsed")
    with c2:
        st.markdown("#### Fichier CA")
        st.caption("N° Dossier · Total des ventes · Client facturation · Etat vente")
        file_ca = st.file_uploader("CA (.xlsx)", type=["xlsx"], key="rd_up_ca", label_visibility="collapsed")

    if not files_pl:
        st.markdown(
            '<ul class="ap-list">'
            '<li><b>1. Plannings + CA</b> <span>— mêmes fichiers que Rentabilité plateau, ou reprise de son calcul.</span></li>'
            '<li><b>2. Calcul</b> <span>— km PTV par trajet, imputation chargé / vide / groupage.</span></li>'
            '<li><b>3. Départements</b> <span>— renta, vide de sortie et d\'entrée, rechargement local.</span></li>'
            '<li><b>4. Contrôles</b> <span>— départements peu représentatifs, repositionnements incomplets.</span></li>'
            '</ul>', unsafe_allow_html=True)
        st.stop()

    stops, erreurs, doublons = P.charger_plannings(files_pl)
    for e in erreurs:
        st.error(f"Lecture impossible — {e}")
    if stops.empty:
        st.stop()
    ca_dict = {}
    if file_ca:
        try:
            ca_dict = P.parse_ca(file_ca)
        except Exception as e:
            st.error(f"Fichier CA : {e}")

    ui.section("Périmètre")
    p1, p2, p3 = st.columns([1.2, 1, 1])
    with p1:
        dmin, dmax = stops["ts_plan"].min(), stops["ts_plan"].max()
        periode = (st.date_input("Période (date planifiée de l'arrêt)", value=(dmin.date(), dmax.date()),
                                 min_value=dmin.date(), max_value=dmax.date(), format="DD/MM/YYYY")
                   if pd.notna(dmin) and pd.notna(dmax) else None)
    with p2:
        types = sorted(t for t in stops["type_trans"].unique() if t)
        f_types = st.multiselect("Type de transport", types, default=[t for t in types if t == "PLA"] or types)
    with p3:
        convention = st.radio("Km à vide imputés au", list(P.CONVENTIONS_VIDE), format_func=P.CONVENTIONS_VIDE.get,
                              help="Sert à la colonne « Renta (convention) ». Le vide de sortie et d'entrée "
                                   "ne dépend pas de ce choix.")

    sel = stops.copy()
    if periode and isinstance(periode, (tuple, list)) and len(periode) == 2:
        a, b = pd.Timestamp(periode[0]), pd.Timestamp(periode[1]) + pd.Timedelta(days=1)
        sel = sel[(sel["ts_plan"] >= a) & (sel["ts_plan"] < b) | sel["ts_plan"].isna()]
    if f_types:
        sel = sel[sel["type_trans"].isin(f_types) | (sel["type_trans"] == "")]
    sel = sel[sel["realise"] & (sel["tracteur"] != "")]

    k = st.columns(4)
    k[0].metric("Tracteurs", int((~sel.groupby("tracteur")["affrete"].first()).sum()) if not sel.empty else 0)
    k[1].metric("Affrétés", int(sel.groupby("tracteur")["affrete"].first().sum()) if not sel.empty else 0)
    k[2].metric("Dossiers", sel["dossier"].nunique())
    k[3].metric("Doublons retirés", doublons)
    if not file_ca:
        st.info("Ajoutez le fichier CA pour obtenir la rentabilité. Sans lui, seuls les km sont calculés.")
    if sel.empty:
        st.warning("Aucun arrêt réalisé dans le périmètre choisi.")
        st.stop()

    if not P.PTV_API_KEY or P.PTV_API_KEY == "METS_TA_CLE_ICI":
        st.error("Clé PTV_API_KEY non configurée : le calcul des distances ne fonctionnera pas.")

    if st.button("Calculer", type="primary"):
        st.session_state.pop("rd_calc", None)
        bar, txt = st.progress(0.0), st.empty()

        def _p(msg, pct):
            txt.caption(msg)
            bar.progress(min(max(pct, 0.0), 1.0))
        try:
            df_traj, dt, _flotte, ctrl_pl = P.calculer(sel, ca_dict, convention, True, progress_cb=_p)
            st.session_state["rd_calc"] = dict(
                sel=sel, df_traj=df_traj, dt=dt, ctrl=ctrl_pl, convention=convention,
                genere=pd.Timestamp.now(tz="Europe/Brussels").strftime("%d/%m/%Y %H:%M"), source="Fichiers")
            bar.progress(1.0)
            txt.empty()
        except Exception as e:
            st.error(f"Erreur pendant le calcul : {e}")
    calcul = st.session_state.get("rd_calc")
    if not calcul:
        st.stop()

# ══════════════════════════════════════════════════════════════════
#  RÉGLAGES D'ANALYSE (sans recalcul PTV)
# ══════════════════════════════════════════════════════════════════

ui.section("Analyse")
r1, r2, r3 = st.columns(3)
with r1:
    seuil_local = st.slider("Rechargement local : à moins de (km)", 10, 150, 50, 5,
                            help="Repositionnement compté comme local s'il fait au plus cette distance.")
with r2:
    min_liv = st.slider("Département représentatif à partir de (livraisons)", 1, 20, 5)
with r3:
    incl_aff = st.toggle("Inclure les affrétés", value=False,
                         help="Par défaut, les affrétés (AFF-xxx) sont comptés à part : leurs km à vide "
                              "ne sont pas un coût pour la flotte.")

R = E.analyser(calcul["sel"], calcul["df_traj"], calcul["dt"], seuil_local, min_liv, incl_aff)
zones, rep, tot = R["zones"], R["repositionnements"], R["totaux"]
ctrl = R["controles"].assign(tracteur="")
cp = calcul["ctrl"]
if cp is not None and not cp.empty:
    ctrl = pd.concat([ctrl, cp.assign(zone="")[["type", "zone", "tracteur", "dossier", "detail"]]], ignore_index=True)

conv_lbl = P.CONVENTIONS_VIDE.get(calcul["convention"], calcul["convention"])
d0, d1 = calcul["sel"]["ts_plan"].min(), calcul["sel"]["ts_plan"].max()
st.caption(f"Source : {calcul['source']} · {d0:%d/%m/%Y} → {d1:%d/%m/%Y} · vide imputé au {conv_lbl.lower()} "
           f"· calculé le {calcul['genere']}")

# ══════════════════════════════════════════════════════════════════
#  RÉSULTATS
# ══════════════════════════════════════════════════════════════════

ui.section("France")
def _n(v, fmt, suffixe=""):
    return (format(v, fmt).replace(",", "\u202f") + suffixe) if pd.notna(v) else "—"


tuiles = [
    (_n(tot["livraisons"], ",.0f"), "Livraisons", f"{int(tot['sorties'])} repositionnements complets"),
    (_n(tot["ca"], ",.0f", " €"), "CA", ""),
    (_n(tot["renta_chargee"], ".2f", " €/km"), "Renta chargée", "CA ÷ km chargés"),
    (_n(tot["renta_conv"], ".2f", " €/km"), "Renta avec vide", conv_lbl),
    (_n(tot["vide_sortie_moy"], ".0f", " km"), "Vide de sortie moyen",
     f"{_n(tot['pct_local'], '.0f', ' %')} rechargés à moins de {seuil_local} km"),
]
st.markdown('<div class="rd-kpis">' + "".join(
    f'<div class="rd-kpi"><div class="v">{v}</div><div class="l">{l}</div><div class="s">{s_}</div></div>'
    for v, l, s_ in tuiles) + "</div>", unsafe_allow_html=True)

fr = zones[zones["france"]].copy()
hors = zones[~zones["france"]].copy()

tabs = st.tabs(["Par département", "Repositionnements", "Hors France", "Contrôles"])

with tabs[0]:
    a1, a2 = st.columns([1, 1])
    with a1:
        vue = st.segmented_control("Vue", ["Renta", "Km à vide", "Tout"], default="Renta",
                                   label_visibility="collapsed") or "Renta"
    with a2:
        rep_seuls = st.toggle(f"Seulement les départements avec au moins {min_liv} livraisons", value=False)
    if rep_seuls:
        fr = fr[fr["representatif"]]
    livres = fr[fr["livraisons"] > 0]
    base = ["libelle", "livraisons"]
    cols = {
        "Renta": base + ["pct_groupage", "ca", "pct_ca", "km_charge", "km_ch_moy", "renta_chargee", "renta_conv",
                         "renta_sortie", "vide_sortie_moy", "livraisons_aff", "representatif"],
        "Km à vide": base + ["sorties", "vide_sortie_moy", "km_sortie_med", "pct_local", "pct_meme_zone",
                             "delai_h", "chargements", "entrees", "vide_entree_moy", "representatif"],
    }.get(vue, list(COLS_ZONES))
    d = (fr if vue == "Km à vide" else livres)[[c for c in cols if c in fr.columns]]
    st.dataframe(styler(d, tot), hide_index=True, width="stretch", height=min(38 * (len(d) + 1) + 3, 720),
                 column_config=_col_config(d.columns))
    st.caption(
        "Couleurs : écart à la moyenne France (vert = mieux). Renta en ratio des sommes : CA ÷ km. "
        "Vide de sortie = km à vide entre un déchargement dans le département et le chargement suivant "
        "(repositionnements complets, un par camion, quel que soit le nombre de dossiers livrés). "
        "Vide d'entrée = km à vide pour venir charger dans le département. "
        "Renta livraison + sortie = CA ÷ (km chargés + repositionnements qui partent du département).")

with tabs[1]:
    if rep.empty:
        st.info("Aucun trajet à vide dans le calcul.")
    else:
        f1, f2 = st.columns([1, 1])
        with f1:
            opts = sorted({z for z in rep["zone_de"].unique() if z} | {z for z in rep["zone_a"].unique() if z},
                          key=lambda z: (not E.est_france(z), z))
            f_z = st.multiselect("Zone (départ ou rechargement)", opts, format_func=E.zone_label, placeholder="Toutes")
        with f2:
            f_st = st.multiselect("Statut", [E.STATUT_COMPLET, E.STATUT_FIN, E.STATUT_DEBUT],
                                  default=[E.STATUT_COMPLET, E.STATUT_FIN])
        r = rep_affichage(rep)
        if f_z:
            r = r[r["zone_de"].isin(f_z) | r["zone_a"].isin(f_z)]
        if f_st:
            r = r[r["statut"].isin(f_st)]
        if not incl_aff:
            r = r[~r["affrete"]]
        r = r.sort_values("km", ascending=False)
        cols = [c for c in COLS_REP if c in r.columns]
        st.dataframe(r[cols], hide_index=True, width="stretch", height=520, column_config=_col_config(cols))
        st.caption(f"{len(r)} repositionnements · {r['km'].sum():,.0f} km à vide".replace(",", " "))

with tabs[2]:
    if hors.empty:
        st.info("Aucun arrêt hors de France.")
    else:
        cols = ["libelle", "livraisons", "ca", "renta_chargee", "renta_conv", "sorties", "vide_sortie_moy",
                "pct_local", "chargements", "entrees", "vide_entree_moy"]
        st.dataframe(styler(hors[cols], tot), hide_index=True, width="stretch",
                     column_config=_col_config(cols))
        st.caption("Pays de chargement surtout : le vide d'entrée dit combien coûte d'aller y chercher du fret.")

with tabs[3]:
    if ctrl.empty:
        st.success("Aucun point à vérifier.")
    else:
        st.dataframe(ctrl[list(COLS_CTRL)].rename(columns=COLS_CTRL), hide_index=True, width="stretch",
                     height=min(38 * (len(ctrl) + 1) + 3, 520))

meta = {
    "Source": calcul["source"], "Période": f"{d0:%d/%m/%Y} → {d1:%d/%m/%Y}",
    "Km à vide imputés au": conv_lbl, "Rechargement local": f"≤ {seuil_local} km",
    "Département représentatif": f"≥ {min_liv} livraisons",
    "Affrétés": "inclus" if incl_aff else "comptés à part", "Calculé le": calcul["genere"],
}
st.write("")
b1, b2 = st.columns([1, 3])
with b1:
    st.download_button("Rapport Excel", data=export_excel(zones, rep, ctrl, meta),
                       file_name="Renta_Departements.xlsx", type="primary", width="stretch",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
with b2:
    if calcul["source"] == "Fichiers" and st.button("Nouveau calcul"):
        st.session_state.pop("rd_calc", None)
        st.rerun()
