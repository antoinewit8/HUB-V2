"""
renta_plateau.py
──────────────────────────────────────────────────────────────────
Outil Rentabilité Plateau
──────────────────────────────────────────────────────────────────
Entrées :
  • Exports planning « Traject_xxxxx.xlsx » (un ou plusieurs) — une ligne par
    arrêt (CH / DECH), dans l'ordre chronologique (colonne Séq.), avec
    N° dossier, Exécution (tracteur · chauffeur · plateau), Km Vehicle…
  • Fichier CA (.xlsx) — N° Dossier | Total des ventes | Client facturation | Etat vente

Logique :
  1. Les lignes sont regroupées par TRACTEUR (lu dans la colonne Exécution),
     tous fichiers confondus, puis remises dans l'ordre (Séq. dans un même
     fichier, date planifiée entre fichiers). Doublons entre fichiers retirés.
  2. On suit la cargaison à bord d'arrêt en arrêt : un CH ajoute le dossier,
     son dernier DECH le retire. Chaque trajet entre deux arrêts est donc :
       – à vide (rien à bord)        → imputé selon la convention choisie ;
       – chargé (un dossier à bord)  → imputé à ce dossier ;
       – groupage (plusieurs)        → parts égales (provisoire, l'outil
                                        groupage fera la vraie répartition).
  3. Km de chaque trajet via PTV. Le compteur (Km Vehicle) sert de contrôle
     global par camion : il n'est pas fiable trajet par trajet (relevés groupés).
  4. Dossiers coupés par la période : le chargement (ou le déchargement)
     manquant est reconstitué depuis Origine / Destination du dossier.
  5. CA joint sur le N° dossier. Un dossier partagé entre deux tracteurs
     (accroche / décroche) voit son CA réparti au prorata des km chargés.
──────────────────────────────────────────────────────────────────
"""

import heapq
import io
import json
import math
import os
import re
import time

import numpy as np
import pandas as pd
import requests
import streamlit as st
import streamlit.components.v1 as components
from dotenv import load_dotenv
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv()

# ══════════════════════════════════════════════════════════════════
#  CONFIG
# ══════════════════════════════════════════════════════════════════

PTV_API_KEY  = os.environ.get("PTV_API_KEY", "METS_TA_CLE_ICI")
PTV_BASE_URL = "https://api.myptv.com/routing/v1"
GEOCODE_URL  = "https://api.myptv.com/geocoding/v1"
HEADERS      = {"apiKey": PTV_API_KEY}
MAX_RETRIES  = 3
RETRY_DELAY  = 2
VEHICLE      = "EUR_TRAILER_TRUCK"

RESULT_SCHEMA      = "plateau-v4"   # change → les résultats en session sont purgés
SEUIL_ALERTE_VIDE  = 150            # km — trajet à vide signalé au-delà
SEUIL_ECART_ODO    = 15             # % — écart compteur / PTV signalé au-delà (par camion)
SEUIL_TEMPS_SITE   = 4.0            # h — temps sur site signalé au-delà

CONVENTIONS_VIDE = {
    "precedent": "Dossier précédent (comme Renta Benne)",
    "suivant":   "Dossier suivant",
    "moitie":    "50/50 entre les deux",
}

PAYS_CODE_TO_ISO2 = {
    "F": "FR", "B": "BE", "D": "DE", "L": "LU", "NL": "NL", "E": "ES", "I": "IT",
    "CH": "CH", "GB": "GB", "A": "AT", "P": "PT", "PL": "PL", "CZ": "CZ",
    "FR": "FR", "BE": "BE", "DE": "DE", "LU": "LU",
}
PAYS_CODE_TO_FULL = {
    "F": "France", "B": "Belgium", "D": "Germany", "L": "Luxembourg", "NL": "Netherlands",
    "E": "Spain", "I": "Italy", "CH": "Switzerland", "GB": "United Kingdom", "A": "Austria",
    "P": "Portugal", "PL": "Poland", "CZ": "Czech Republic",
    "FR": "France", "BE": "Belgium", "DE": "Germany", "LU": "Luxembourg",
}

# Couleurs (cohérentes avec le hub)
C_CHARGE, C_VIDE, C_GROUP, C_RECON = "#0071e3", "#8e8e93", "#ff9500", "#af52de"


# ══════════════════════════════════════════════════════════════════
#  UTILITAIRES
# ══════════════════════════════════════════════════════════════════

def _norm(s) -> str:
    s = str(s).lower()
    for a, b in [("é", "e"), ("è", "e"), ("ê", "e"), ("à", "a"), ("â", "a"), ("ô", "o"),
                 ("û", "u"), ("î", "i"), ("ù", "u"), ("ç", "c"), ("°", "")]:
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]", "", s)


def _clean(v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    v = str(v).strip()
    return "" if v.lower() in ("nan", "none", "nat") else v


def _cp_norm(cp) -> str:
    """'07470' / '7470.0' / 'F 7470' → '7470' (comparaison seulement)."""
    s = re.sub(r"\.0+$", "", _clean(cp)).upper()
    s = re.sub(r"^(F|B|D|L|NL|E|I|CH|A|P|FR|BE|DE|LU)[\s\-]*(?=\d)", "", s)
    s = re.sub(r"[^A-Z0-9]", "", s)
    return s.lstrip("0") or s


def _to_float(v):
    s = _clean(v).replace("\xa0", "").replace(" ", "").replace("€", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan


def _to_ts(date_v, heure_v=None):
    """Date (datetime ou texte) + heure ('07:33:00') → Timestamp, NaT si vide."""
    d = pd.to_datetime(_clean(date_v) or None, errors="coerce", dayfirst=False)
    if pd.isna(d) or d.year < 2000:
        return pd.NaT
    h = _clean(heure_v)
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?", h)
    if m:
        d = d.normalize() + pd.Timedelta(hours=int(m.group(1)), minutes=int(m.group(2)),
                                         seconds=int(m.group(3) or 0))
    return d


def haversine_km(a, b):
    if not a or not b:
        return None
    lat1, lon1, lat2, lon2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = (math.sin((lat2 - lat1) / 2) ** 2
         + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2)
    return 6371.0 * 2 * math.asin(math.sqrt(h))


# ══════════════════════════════════════════════════════════════════
#  PTV — géocodage et distances (mis en cache pour la session)
# ══════════════════════════════════════════════════════════════════

def _ptv_get(url, params):
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.get(url, params=params, headers=HEADERS, timeout=30)
            if resp.status_code == 429:
                time.sleep(RETRY_DELAY * attempt)
                continue
            if resp.status_code != 200:
                if attempt < MAX_RETRIES:
                    time.sleep(RETRY_DELAY)
                    continue
                return None
            return resp.json()
        except Exception:
            time.sleep(RETRY_DELAY)
    return None


def _first_pos(js):
    locs = (js or {}).get("locations", [])
    if locs:
        pos = locs[0]["referencePosition"]
        return (pos["latitude"], pos["longitude"])
    return None


@st.cache_data(show_spinner=False)
def geocode_stop(ville: str, cp: str, pays: str):
    """Priorité : CP + pays → ville, CP, pays → ville, pays → ville seule."""
    iso2 = PAYS_CODE_TO_ISO2.get(pays.upper(), "")
    full = PAYS_CODE_TO_FULL.get(pays.upper(), "")
    if cp and iso2:
        r = _first_pos(_ptv_get(f"{GEOCODE_URL}/locations/by-postal-code",
                                {"postalCode": cp, "countryCode": iso2}))
        if r:
            return r
    for q in ([f"{ville}, {cp}, {full}"] if ville and cp and full else []) + \
             ([f"{ville}, {full}"] if ville and full else []) + \
             ([f"{cp} {ville}"] if ville and cp and not full else []) + ([ville] if ville else []):
        r = _first_pos(_ptv_get(f"{GEOCODE_URL}/locations/by-text", {"searchText": q}))
        if r:
            return r
    return None


@st.cache_data(show_spinner=False)
def route_km(a: tuple, b: tuple):
    """Distance routière PTV (poids lourd) entre deux points GPS, en km."""
    if not a or not b:
        return None
    if haversine_km(a, b) < 0.3:
        return 0.0
    js = _ptv_get(f"{PTV_BASE_URL}/routes",
                  [("profile", VEHICLE), ("waypoints", f"{a[0]},{a[1]}"), ("waypoints", f"{b[0]},{b[1]}")])
    if js is None:
        return None
    return round(js.get("distance", 0) / 1000, 1)


# ══════════════════════════════════════════════════════════════════
#  LECTURE DES EXPORTS PLANNING
# ══════════════════════════════════════════════════════════════════

PLAN_COLS = {
    "dossier":    ["N° dossier", "N° Dossier", "Dossier"],
    "date":       ["Date"],
    "heure":      ["De"],
    "seq":        ["Séq.", "Seq.", "Séq"],
    "activite":   ["Activité"],
    "nom":        ["Nom"],
    "cp":         ["Code postal"],
    "ville":      ["Ville"],
    "vente":      ["Vente"],
    "execution":  ["Exécution"],
    "date_ob":    ["Date OB"],
    "ob_de":      ["OB de"],
    "ob_a":       ["OB à"],
    "odo":        ["Km Vehicle"],
    "poids":      ["Poids"],
    "produit":    ["Product Nature", "Produit"],
    "client_pl":  ["Client"],
    "type_trans": ["Type trans."],
    "statut":     ["Statut dos."],
    "nom_orig":   ["Origine"],
    "pays_orig":  ["Pays d'origine"],
    "cp_orig":    ["CP orig."],
    "ville_orig": ["Ville origine"],
    "nom_dest":   ["Destination"],
    "pays_dest":  ["Pays de destination"],
    "cp_dest":    ["CP dest."],
    "ville_dest": ["Ville destination"],
}
OBLIGATOIRES = ["dossier", "seq", "activite", "cp", "ville", "execution"]


LONGUEUR_EXECUTION = 50     # Traject coupe la colonne Exécution à 50 caractères
SEUIL_RESSOURCE_FICHIER = 0.6   # part des lignes où une ressource doit figurer pour « posséder » le fichier
_RE_AFF = re.compile(r"^AFF-[A-Z0-9]+$", re.I)


def _est_plaque(t: str) -> bool:
    return bool(re.search(r"\d", t)) and len(t) >= 4 and not t.upper().startswith("AFF-")


def decouper_execution(txt: str):
    """
    La colonne Exécution enchaîne une ou plusieurs ressources, parfois suivies
    (ou précédées) du nom de la société affrétée :
      'WM7434 EL-ABBADI QACF329'                          → tracteur propre
      'AFF-CBFR LEDOUX QAJE768 CB FRANCE SAS'             → affrété : code AFF, chauffeur, plateau, société
      'TRANSPORT DE LAIT BOUCHART AFF-TLB LAHAYE QAFY498' → société avant la ressource
      'RJ8636 KRAJNAK MZ4772 CB FRANCE SAS AFF-CBFR LEDOU' → deux ressources (transbordement), texte coupé
    Une ressource = tracteur (plaque ou code AFF-xxx) + chauffeur + plateau.
    Clé : la plaque du tracteur ; pour un affrété, « code AFF + chauffeur »
    (le code AFF est partagé par tous les chauffeurs d'une même société).
    Retourne (ressources, sociétés, fragment) — fragment = début de ressource
    coupé à 50 caractères (ex. 'AFF-CBFR LEDOU'), à comparer aux clés connues.
    """
    s = _clean(txt)
    toks = s.split()
    tronque = len(s) >= LONGUEUR_EXECUTION
    ressources, societes, buf, cur = [], [], [], None

    def _fermer():
        nonlocal cur
        if cur:
            ressources.append(cur)
            cur = None

    for t in toks:
        aff = bool(_RE_AFF.match(t))
        if aff or (_est_plaque(t) and (cur is None or cur["plateau"])):
            _fermer()
            if buf:
                societes.append(" ".join(buf))
                buf = []
            cur = dict(tracteur=t.upper(), chauffeur="", plateau="", affrete=aff)
        elif _est_plaque(t):
            cur["plateau"] = t.upper()
        elif cur is not None and not cur["plateau"]:
            cur["chauffeur"] = f"{cur['chauffeur']} {t}".strip()
        else:
            _fermer()
            buf.append(t)
    _fermer()

    fragment = ""
    if tronque and toks:
        if buf:                                   # 'TRANSPORT DE LAIT BOUCHART A' → 'A' coupé
            fragment = buf.pop()
        elif ressources and not ressources[-1]["plateau"]:
            r = ressources.pop()                  # 'AFF-CBFR LEDOU' → ressource incomplète
            fragment = f"{r['tracteur']} {r['chauffeur']}".strip()
    if buf:
        societes.append(" ".join(buf))
    for r in ressources:
        r["cle"] = f"{r['tracteur']} {r['chauffeur']}".strip() if r["affrete"] else r["tracteur"]
    return ressources, societes, fragment


def affecter_ressources(rows: list):
    """
    Un export Traject = le planning d'UNE ressource (tracteur propre ou chauffeur affrété),
    mais la colonne Exécution liste toutes les ressources du dossier (transbordement,
    accroche / décroche). On cherche la ressource présente sur la majorité des lignes du
    fichier : toutes les lignes lui sont attribuées. Sinon (fichier multi-tracteurs),
    chaque ligne va à la première ressource lue.
    """
    compte = {}
    for r in rows:
        for c in dict.fromkeys(x["cle"] for x in r["_res"]):
            compte[c] = compte.get(c, 0) + 1
    if not compte:
        return
    # les fragments coupés comptent pour la clé qu'ils prolongent
    for r in rows:
        if r["_frag"]:
            for c in compte:
                if c.startswith(r["_frag"]) and not any(x["cle"] == c for x in r["_res"]):
                    compte[c] += 1
    principale = max(compte, key=lambda c: (compte[c], c))
    proprio = compte[principale] >= SEUIL_RESSOURCE_FICHIER * len(rows)

    ref = next(x for r in rows for x in r["_res"] if x["cle"] == principale)
    soc = {}
    for r in rows:
        for s_ in r["_soc"]:
            soc[s_] = soc.get(s_, 0) + 1
    societe = max(soc, key=soc.get) if soc and ref["affrete"] else ""

    for r in rows:
        if proprio:
            res = next((x for x in r["_res"] if x["cle"] == principale), ref)
            autres = [x for x in r["_res"] if x["cle"] != principale]
            r["tracteur"], r["chauffeur"] = principale, res["chauffeur"] or ref["chauffeur"]
            r["plateau"], r["affrete"], r["societe"] = res["plateau"] or ref["plateau"], ref["affrete"], societe
            r["accroche"] = bool(autres)
            r["autres_ressources"] = ", ".join(f"{x['cle']} {x['chauffeur']}".strip() if not x["affrete"]
                                               else x["cle"] for x in autres)
        else:
            res = r["_res"][0] if r["_res"] else None
            r["tracteur"] = res["cle"] if res else ""
            r["chauffeur"], r["plateau"] = (res["chauffeur"], res["plateau"]) if res else ("", "")
            r["affrete"] = bool(res and res["affrete"])
            r["societe"] = (r["_soc"][0] if r["_soc"] else "") if r["affrete"] else ""
            r["accroche"] = len(r["_res"]) > 1
            r["autres_ressources"] = ", ".join(x["cle"] for x in r["_res"][1:])


def _pays_arret(cp, act, r):
    """Le planning ne donne pas le pays de l'arrêt : on le déduit d'Origine / Destination."""
    c = _cp_norm(cp)
    po, pd_ = r.get("pays_orig", "").upper(), r.get("pays_dest", "").upper()
    if c and c == _cp_norm(r.get("cp_orig")) and po:
        return po
    if c and c == _cp_norm(r.get("cp_dest")) and pd_:
        return pd_
    if act not in ("CH", "DECH"):
        # Arrêt au dépôt / reprise (ex. DO à Werbomont) : Origine et Destination
        # sont celles du dossier, pas du lieu → on ne s'y fie pas.
        brut = _clean(cp).upper()
        if re.match(r"^L[\s\-]?\d{4}$", brut):
            return "L"
        guess = ""
    else:
        guess = (po if act == "CH" else pd_) or po or pd_
    digits = re.sub(r"\D", "", _clean(cp))
    if len(digits) == 5 and guess in ("B", "L", "NL", "A", "CH", ""):
        guess = "F"
    elif len(digits) == 4 and guess in ("F", "D", "E", "I", ""):
        guess = "B"
    return guess


def parse_planning(file, file_idx: int):
    df = pd.read_excel(file, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    by_norm = {}
    for c in df.columns:
        by_norm.setdefault(_norm(c), c)
    col_map = {}
    for role, cands in PLAN_COLS.items():
        col_map[role] = next((by_norm[_norm(c)] for c in cands if _norm(c) in by_norm), None)
    manquantes = [r for r in OBLIGATOIRES if not col_map.get(r)]
    if manquantes:
        raise ValueError(f"colonnes introuvables : {manquantes}")

    nom_fichier = getattr(file, "name", f"fichier {file_idx + 1}")
    rows = []
    for i, src in df.iterrows():
        r = {role: (_clean(src[col]) if col else "") for role, col in col_map.items()}
        if not r["dossier"]:
            continue
        r["dossier"] = re.sub(r"\.0+$", "", r["dossier"])
        act = r["activite"].upper()
        r["activite"] = "CH" if act.startswith("CH") else ("DECH" if act.startswith("DE") else act)
        r["_res"], r["_soc"], r["_frag"] = decouper_execution(r["execution"])
        r["pays"] = _pays_arret(r["cp"], r["activite"], r)
        r["cp"] = re.sub(r"\.0+$", "", r["cp"])
        r["seq_num"] = _to_float(r["seq"])
        r["ts_plan"] = _to_ts(src[col_map["date"]] if col_map["date"] else None, r["heure"])
        r["ts_arr"] = _to_ts(src[col_map["date_ob"]], r["ob_de"]) if col_map["date_ob"] else pd.NaT
        r["ts_dep"] = _to_ts(src[col_map["date_ob"]], r["ob_a"]) if col_map["date_ob"] else pd.NaT
        if pd.notna(r["ts_arr"]) and pd.notna(r["ts_dep"]) and r["ts_dep"] < r["ts_arr"]:
            r["ts_dep"] += pd.Timedelta(days=1)          # départ après minuit
        r["realise"] = pd.notna(r["ts_arr"])
        odo = _to_float(r["odo"])
        r["odo_km"] = odo if odo and odo > 0 else np.nan
        r["poids_t"] = _to_float(r["poids"])
        r["fichier"] = nom_fichier
        r["file_idx"] = file_idx
        r["ligne"] = int(i) + 2
        rows.append(r)
    affecter_ressources(rows)
    for r in rows:
        for k in ("_res", "_soc", "_frag"):
            r.pop(k, None)
    df_out = pd.DataFrame(rows)
    if df_out.empty:
        return df_out
    # Sans aucun pointage OB (cas des affrétés) : un arrêt est considéré réalisé
    # dès que sa date planifiée est passée (aujourd'hui inclus).
    df_out["sans_pointage"] = False
    aujourd_hui = pd.Timestamp.now().normalize() + pd.Timedelta(days=1)
    for tr, idx in df_out.groupby("tracteur").groups.items():
        if not df_out.loc[idx, "realise"].any():
            df_out.loc[idx, "sans_pointage"] = True
            df_out.loc[idx, "realise"] = df_out.loc[idx, "ts_plan"].lt(aujourd_hui).fillna(False)
    return df_out


def ordonner(df: pd.DataFrame) -> pd.DataFrame:
    """
    Ordre de la tournée par tracteur.
    Dans un même fichier : la Séq. fait foi. Entre fichiers : fusion sur la date
    planifiée, sans jamais casser l'ordre Séq. d'un fichier.
    """
    out = []
    for tr, g in df.groupby("tracteur", sort=True):
        blocs = []
        for _, b in g.groupby("file_idx", sort=True):
            b = b.sort_values(["seq_num", "ligne"], kind="stable")
            blocs.append([(r.ts_plan if pd.notna(r.ts_plan) else pd.Timestamp.max, r.file_idx, k, idx)
                          for k, (idx, r) in enumerate(b.iterrows())])
        ordre = [t[3] for t in heapq.merge(*blocs, key=lambda t: (t[0], t[1], t[2]))]
        o = g.loc[ordre].copy()
        o["ordre"] = range(1, len(o) + 1)
        out.append(o)
    return pd.concat(out, ignore_index=True) if out else df


def charger_plannings(files):
    frames, erreurs = [], []
    for k, f in enumerate(files):
        try:
            frames.append(parse_planning(f, k))
        except Exception as e:
            erreurs.append(f"{getattr(f, 'name', k)} : {e}")
    if not frames:
        return pd.DataFrame(), erreurs, 0
    df = pd.concat(frames, ignore_index=True)
    n0 = len(df)
    df = df.drop_duplicates(subset=["tracteur", "dossier", "activite", "cp", "ville", "seq", "ts_plan"],
                            keep="first")
    return attribuer_roles(ordonner(df.reset_index(drop=True))), erreurs, n0 - len(df)


def attribuer_roles(df: pd.DataFrame) -> pd.DataFrame:
    """
    Rôle de chaque arrêt pour la cargaison à bord : CH (monte), DECH (descend), AUTRE.
    Une activité autre que CH / DECH (ex. DO au dépôt de Werbomont) est une reprise
    si le dossier n'est pas encore à bord et revient plus loin (le camion le récupère),
    une dépose s'il est déjà à bord. Après le dernier déchargement du dossier
    (ex. DCR = décrochage), elle reste neutre.
    """
    if df.empty:
        return df
    df["role"] = df["activite"].where(df["activite"].isin(["CH", "DECH"]), "AUTRE")
    for tr, g in df.groupby("tracteur", sort=False):
        g = g.sort_values("ordre")
        vus, livres = set(), set()
        idx, dos, act = list(g.index), list(g["dossier"]), list(g["activite"])
        for k, (i, d, a) in enumerate(zip(idx, dos, act)):
            if a == "CH":
                vus.add(d)
            elif a == "DECH":
                livres.add(d)
            elif a not in ("CH", "DECH", "LAV") and d:
                if d in livres and d not in dos[k + 1:]:
                    pass                      # ex. DCR après le déchargement : rien à bord, arrêt neutre
                elif d in vus:
                    df.at[i, "role"] = "DECH"
                elif any(d2 == d and a2 not in ("CH", "LAV") for d2, a2 in zip(dos[k + 1:], act[k + 1:])):
                    # Reprise seulement si le dossier est réellement déchargé / déposé plus loin.
                    # Un accrochage (ACR) suivi d'un simple lavage = citerne vide qu'on emmène laver.
                    df.at[i, "role"] = "CH"
                    vus.add(d)
    return df


# ══════════════════════════════════════════════════════════════════
#  FICHIER CA
# ══════════════════════════════════════════════════════════════════

def parse_ca(file):
    df = pd.read_excel(file, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    by_norm = {_norm(c): c for c in df.columns}

    def _find(cands):
        return next((by_norm[_norm(c)] for c in cands if _norm(c) in by_norm), None)

    c_dos = _find(["N° Dossier", "N°Dossier", "Dossier"])
    c_tot = _find(["Total des ventes", "Total ventes", "Total des vente"])
    c_cli = _find(["Client facturation", "Client"])
    c_eta = _find(["Etat vente", "État vente", "Etat"])
    if not c_dos or not c_tot:
        raise ValueError(f"colonnes CA manquantes — disponibles : {list(df.columns)}")
    ca = {}
    for _, r in df.iterrows():
        d = re.sub(r"\.0+$", "", _clean(r[c_dos]))
        if not d:
            continue
        v = _to_float(r[c_tot])
        v = 0.0 if np.isnan(v) else v
        if d in ca:
            ca[d]["ca"] += v
        else:
            ca[d] = {"ca": v, "client": _clean(r[c_cli]) if c_cli else "",
                     "etat": _clean(r[c_eta]) if c_eta else ""}
    return ca


# ══════════════════════════════════════════════════════════════════
#  MOTEUR : cargaison à bord, trajets, imputation
# ══════════════════════════════════════════════════════════════════

def _lieu(r, prefix=""):
    if prefix:
        return (r.get(f"ville_{prefix}", ""), _clean(r.get(f"cp_{prefix}", "")).split(".")[0],
                r.get(f"pays_{prefix}", "").upper())
    return (r["ville"], r["cp"], r["pays"])


def construire_trajets(stops: pd.DataFrame, reconstituer: bool = True):
    """
    Retourne la liste des trajets (sans km) avec leur cargaison et le type.
    Chaque trajet : tracteur, de/à (lieu + arrêt), dossiers à bord, type, dossier_prec / dossier_suiv.
    """
    trajets, controles = [], []
    for tr, g in stops.groupby("tracteur", sort=True):
        g = g.sort_values("ordre").to_dict("records")
        dernier_dech = {}
        premier = {}
        idx_ch, idx_dech = {}, {}
        for k, r in enumerate(g):
            premier.setdefault(r["dossier"], (k, r["role"]))
            if r["role"] == "DECH":
                dernier_dech[r["dossier"]] = k
                idx_dech.setdefault(r["dossier"], []).append(k)
            elif r["role"] == "CH":
                idx_ch.setdefault(r["dossier"], []).append(k)

        # Ordre CH / DECH incohérent : le dossier a bien ses deux bouts dans le fichier,
        # mais un CH arrive après son dernier DECH (activité validée en retard sur l'OBC,
        # Séq. inversée…). Sans correction, le dossier serait vu « chargé avant » ET
        # « déchargé après » la période et resterait à bord de toute la tournée.
        # → à bord uniquement entre son premier et son dernier arrêt CH/DECH.
        inverses = {}
        for d in set(idx_ch) & set(idx_dech):
            if max(idx_ch[d]) > max(idx_dech[d]):
                ks = idx_ch[d] + idx_dech[d]
                inverses[d] = (min(ks), max(ks))
        for d, (k1, k2) in sorted(inverses.items()):
            r1, r2 = g[k1], g[k2]
            controles.append(dict(
                tracteur=tr, dossier=d, type="Ordre CH/DECH incohérent",
                detail=f"{r1['activite']} {r1['ville']} (ordre {r1['ordre']}) avant "
                       f"{r2['activite']} {r2['ville']} (ordre {r2['ordre']}) — dossier compté à bord "
                       f"uniquement entre ces deux arrêts. Vérifier l'horodatage dans Traject."))

        # Dossiers déjà à bord au début (chargés avant la période)
        a_bord = {d for d, (k, act) in premier.items() if act == "DECH" and d not in inverses}
        for d in sorted(a_bord):
            controles.append(dict(tracteur=tr, dossier=d, type="Chargement avant la période",
                                  detail="Dossier déchargé sans chargement dans le fichier — "
                                         + ("chargement reconstitué depuis Origine." if reconstituer
                                            else "km chargés incomplets.")))
        if reconstituer and g and g[0]["role"] == "DECH" and g[0]["dossier"] in a_bord:
            r0 = g[0]
            o = _lieu(r0, "orig")
            if o[0] or o[1]:
                trajets.append(dict(
                    tracteur=tr, chauffeur=r0["chauffeur"], ordre=0, reconstitue=True,
                    de_nom=r0.get("nom_orig", ""), de=o, de_act="CH*", de_ts=pd.NaT, de_odo=np.nan,
                    a_nom=r0["nom"], a=_lieu(r0), a_act=r0["activite"], a_ts=r0["ts_arr"], a_odo=r0["odo_km"],
                    a_bord=tuple(sorted(a_bord)), dossier_prec="", dossier_suiv=r0["dossier"]))

        dernier_livre = ""
        for k in range(len(g)):
            r = g[k]
            d = r["dossier"]
            if d in inverses:
                if k == inverses[d][0]:
                    a_bord.add(d)
                elif k == inverses[d][1]:
                    a_bord.discard(d)
                    dernier_livre = d
            elif r["role"] == "CH":
                a_bord.add(d)
            elif r["role"] == "DECH" and dernier_dech.get(d) == k:
                a_bord.discard(d)
                dernier_livre = d
            if k == len(g) - 1:
                break
            n = g[k + 1]
            suiv = next((x["dossier"] for x in g[k + 1:] if x["role"] == "CH"), "")
            trajets.append(dict(
                tracteur=tr, chauffeur=r["chauffeur"] or n["chauffeur"], ordre=k + 1, reconstitue=False,
                de_nom=r["nom"], de=_lieu(r), de_act=r["activite"], de_ts=r["ts_dep"], de_odo=r["odo_km"],
                a_nom=n["nom"], a=_lieu(n), a_act=n["activite"], a_ts=n["ts_arr"], a_odo=n["odo_km"],
                a_bord=tuple(sorted(a_bord)), dossier_prec=dernier_livre, dossier_suiv=suiv))

        # Dossiers encore à bord à la fin (déchargement après la période / pas encore réalisé)
        if a_bord:
            for d in sorted(a_bord):
                controles.append(dict(tracteur=tr, dossier=d, type="Déchargement après la période",
                                      detail="Dossier chargé sans déchargement dans le fichier — "
                                             + ("déchargement reconstitué depuis Destination." if reconstituer
                                                else "km chargés incomplets.")))
            if reconstituer:
                rl = g[-1]
                # Un seul trajet de fin, vers la destination du dossier le plus récemment chargé
                d_fin = next((x["dossier"] for x in reversed(g) if x["dossier"] in a_bord), sorted(a_bord)[0])
                ref = next(x for x in g if x["dossier"] == d_fin)
                dst = _lieu(ref, "dest")
                if dst[0] or dst[1]:
                    trajets.append(dict(
                        tracteur=tr, chauffeur=rl["chauffeur"], ordre=len(g), reconstitue=True,
                        de_nom=rl["nom"], de=_lieu(rl), de_act=rl["activite"], de_ts=rl["ts_dep"],
                        de_odo=rl["odo_km"], a_nom=ref.get("nom_dest", ""), a=dst, a_act="DECH*",
                        a_ts=pd.NaT, a_odo=np.nan, a_bord=tuple(sorted(a_bord)),
                        dossier_prec=dernier_livre, dossier_suiv=""))

    for t in trajets:
        nb = len(t["a_bord"])
        t["type"] = "Reconstitué" if t["reconstitue"] else ("À vide" if nb == 0 else
                                                            ("Chargé" if nb == 1 else "Groupage"))
    return trajets, controles


def imputer(trajets, convention: str):
    """Part de km de chaque trajet attribuée à chaque dossier → liste (tracteur, dossier, km_ch, km_vide)."""
    parts = []
    for t in trajets:
        km = t.get("km")
        if km is None or (isinstance(km, float) and np.isnan(km)):
            continue
        if t["a_bord"]:
            share = km / len(t["a_bord"])
            for d in t["a_bord"]:
                parts.append((t["tracteur"], d, share, 0.0))
        else:
            p, s = t["dossier_prec"], t["dossier_suiv"]
            if convention == "suivant":
                cibles = [s or p]
            elif convention == "moitie" and p and s and p != s:
                cibles = [p, s]
            else:
                cibles = [p or s]
            cibles = [c for c in cibles if c]
            for c in cibles:
                parts.append((t["tracteur"], c, 0.0, km / len(cibles)))
    return pd.DataFrame(parts, columns=["tracteur", "dossier", "km_charge", "km_vide"])


def calculer(stops, ca_dict, convention, reconstituer, progress_cb=None, geocoder=None, distance=None):
    geocoder = geocoder or geocode_stop
    distance = distance or route_km
    trajets, controles = construire_trajets(stops, reconstituer)

    # 1. Géocodage des lieux uniques
    lieux = sorted({t["de"] for t in trajets} | {t["a"] for t in trajets}
                   | {_lieu(r) for r in stops.to_dict("records")})
    coords = {}
    for i, l in enumerate(lieux):
        if progress_cb:
            progress_cb(f"Géocodage {i + 1}/{len(lieux)} : {l[0]} {l[1]}", i / max(len(lieux), 1) * 0.3)
        coords[l] = geocoder(*l)
        if coords[l] is None and (l[0] or l[1]):
            controles.append(dict(tracteur="", dossier="", type="Géocodage échoué",
                                  detail=f"{l[0]} {l[1]} ({l[2]}) — km des trajets concernés non calculés."))

    # 2. Km PTV par trajet
    for i, t in enumerate(trajets):
        if progress_cb:
            progress_cb(f"Km {i + 1}/{len(trajets)} : {t['de'][0]} → {t['a'][0]}",
                        0.3 + i / max(len(trajets), 1) * 0.7)
        a, b = coords.get(t["de"]), coords.get(t["a"])
        t["km"] = 0.0 if t["de"] == t["a"] else distance(a, b)
        t["lat_de"], t["lon_de"] = (a or (np.nan, np.nan))
        t["lat_a"], t["lon_a"] = (b or (np.nan, np.nan))
        odo = t["a_odo"] - t["de_odo"] if pd.notna(t["a_odo"]) and pd.notna(t["de_odo"]) else np.nan
        t["km_compteur"] = odo
        t["releve_groupe"] = bool(pd.notna(odo) and t["km"] and t["km"] > 20 and odo < 0.3 * t["km"])
        if pd.notna(odo) and (odo < 0 or (t["km"] is not None and not np.isnan(t["km"])
                                          and odo > 2 * t["km"] + 100)):
            controles.append(dict(tracteur=t["tracteur"], dossier=", ".join(t["a_bord"]),
                                  type="Relevé compteur incohérent",
                                  detail=f"{t['de'][0]} → {t['a'][0]} : compteur {odo:+.0f} km vs PTV "
                                         f"{t['km']:.0f} km — activité probablement validée hors séquence."))
        if t["type"] == "À vide" and t["km"] and t["km"] > SEUIL_ALERTE_VIDE:
            controles.append(dict(tracteur=t["tracteur"], dossier=t["dossier_suiv"] or t["dossier_prec"],
                                  type=f"Vide > {SEUIL_ALERTE_VIDE} km",
                                  detail=f"{t['de'][0]} → {t['a'][0]} : {t['km']:.0f} km à vide."))

    df_traj = pd.DataFrame(trajets)

    # 3. Imputation par dossier / tracteur
    parts = imputer(trajets, convention)
    dt = parts.groupby(["tracteur", "dossier"], as_index=False)[["km_charge", "km_vide"]].sum()

    # Dossiers présents dans le planning mais sans aucun km imputé (garder pour le CA)
    presents = stops.groupby(["tracteur", "dossier"], as_index=False).size()[["tracteur", "dossier"]]
    dt = presents.merge(dt, on=["tracteur", "dossier"], how="left").fillna({"km_charge": 0.0, "km_vide": 0.0})

    # 4. CA (réparti au prorata des km chargés si le dossier a roulé sur plusieurs tracteurs)
    tot_ch = dt.groupby("dossier")["km_charge"].transform("sum")
    nb_tr = dt.groupby("dossier")["tracteur"].transform("count")
    dt["part_ca"] = np.where(tot_ch > 0, dt["km_charge"] / tot_ch.replace(0, np.nan), 1 / nb_tr)
    dt["ca_dossier"] = dt["dossier"].map(lambda d: ca_dict.get(d, {}).get("ca", np.nan))
    dt["ca_trouve"] = dt["dossier"].map(lambda d: d in ca_dict)
    dt["ca"] = (dt["ca_dossier"].fillna(0) * dt["part_ca"]).round(2)
    dt["client"] = dt["dossier"].map(lambda d: ca_dict.get(d, {}).get("client", ""))
    dt["etat_vente"] = dt["dossier"].map(lambda d: ca_dict.get(d, {}).get("etat", ""))
    dt["km_total"] = dt["km_charge"] + dt["km_vide"]
    dt["renta"] = (dt["ca"] / dt["km_total"].replace(0, np.nan)).round(3).where(dt["ca_trouve"])

    # Infos dossier depuis le planning
    info = stops.sort_values("ordre").groupby(["tracteur", "dossier"]).agg(
        chauffeur=("chauffeur", "first"),
        client_planning=("client_pl", "first"),
        produit=("produit", "first"),
        poids=("poids_t", "max"),
        vente_planning=("vente", "first"),
        chargement=("ville", lambda s: " · ".join(dict.fromkeys(
            stops.loc[s.index][stops.loc[s.index, "role"] == "CH"]["ville"]))),
        dechargement=("ville", lambda s: " · ".join(dict.fromkeys(
            stops.loc[s.index][stops.loc[s.index, "role"] == "DECH"]["ville"]))),
        nb_ch=("role", lambda s: int((s == "CH").sum())),
        nb_dech=("role", lambda s: int((s == "DECH").sum())),
        date_debut=("ts_plan", "min"),
        premier_arret=("ordre", "min"),
        realise=("realise", "all"),
    ).reset_index()
    dt = dt.merge(info, on=["tracteur", "dossier"], how="left")
    # Chargement / déchargement hors période : on affiche l'origine / la destination du dossier (avec *)
    orig = stops.groupby("dossier")["ville_orig"].first()
    dest = stops.groupby("dossier")["ville_dest"].first()
    vide_ch = dt["chargement"].fillna("") == ""
    dt.loc[vide_ch, "chargement"] = dt.loc[vide_ch, "dossier"].map(orig).fillna("").map(lambda v: f"{v} *" if v else "")
    vide_de = dt["dechargement"].fillna("") == ""
    dt.loc[vide_de, "dechargement"] = dt.loc[vide_de, "dossier"].map(dest).fillna("").map(lambda v: f"{v} *" if v else "")
    dt["client"] = np.where(dt["client"] != "", dt["client"], dt["client_planning"])
    groupes = set()
    for t in trajets:
        if len(t["a_bord"]) > 1:
            groupes.update((t["tracteur"], d) for d in t["a_bord"])
    dt["groupage"] = [(r.tracteur, r.dossier) in groupes for r in dt.itertuples()]

    for r in dt[~dt["ca_trouve"]].itertuples():
        controles.append(dict(tracteur=r.tracteur, dossier=r.dossier, type="Dossier sans CA",
                              detail=f"Absent du fichier CA (vente planning : {r.vente_planning or '—'})."))

    # 5. Synthèse par tracteur
    df_flotte = synthese_flotte(stops, df_traj, dt, controles)
    df_ctrl = pd.DataFrame(controles, columns=["tracteur", "dossier", "type", "detail"])
    return df_traj, dt, df_flotte, df_ctrl


def synthese_flotte(stops, df_traj, dt, controles):
    lignes = []
    for tr, g in stops.groupby("tracteur", sort=True):
        tj = df_traj[df_traj["tracteur"] == tr] if not df_traj.empty else df_traj
        km = tj["km"].fillna(0) if not tj.empty else pd.Series(dtype=float)
        km_tot = float(km.sum())
        km_vide = float(km[tj["type"] == "À vide"].sum()) if not tj.empty else 0.0
        km_group = float(km[tj["type"] == "Groupage"].sum()) if not tj.empty else 0.0
        ca = float(dt.loc[dt["tracteur"] == tr, "ca"].sum())
        jours = g["ts_plan"].dropna().dt.normalize().nunique()

        # Contrôle compteur : du premier au dernier arrêt réalisé avec un relevé
        km_cpt, ecart = np.nan, np.nan
        go = g.sort_values("ordre")
        odo = go[go["odo_km"].notna()]
        if len(odo) >= 2 and not tj.empty:
            o1, o2 = odo.iloc[0]["ordre"], odo.iloc[-1]["ordre"]
            km_cpt = float(odo.iloc[-1]["odo_km"] - odo.iloc[0]["odo_km"])
            km_ptv_span = float(tj[(~tj["reconstitue"]) & (tj["ordre"] >= o1) & (tj["ordre"] < o2)]["km"]
                                .fillna(0).sum())
            if km_ptv_span > 0:
                ecart = round((km_cpt - km_ptv_span) / km_ptv_span * 100, 1)
                if abs(ecart) > SEUIL_ECART_ODO:
                    controles.append(dict(tracteur=tr, dossier="", type="Écart compteur / PTV",
                                          detail=f"Compteur {km_cpt:,.0f} km vs PTV {km_ptv_span:,.0f} km "
                                                 f"sur les arrêts réalisés ({ecart:+.1f} %)."))
        chauffeurs = list(dict.fromkeys(c for c in g.sort_values("ordre")["chauffeur"] if c))
        plateaux = list(dict.fromkeys(p for p in g.sort_values("ordre")["plateau"] if p))
        if len(plateaux) > 1:
            controles.append(dict(tracteur=tr, dossier="", type="Accroche / décroche",
                                  detail=f"Plateaux : {', '.join(plateaux)} · "
                                         f"chauffeurs : {', '.join(chauffeurs) or '—'}."))
        partages = g[g["accroche"]].groupby("dossier")["autres_ressources"].first()
        for d, autres in partages.items():
            controles.append(dict(tracteur=tr, dossier=d, type="Dossier partagé (transbordement)",
                                  detail=f"Aussi exécuté par {autres or 'une autre ressource'}. Sans le planning "
                                         f"de cette ressource, tout le CA du dossier reste sur {tr}."))
        if g["sans_pointage"].any():
            controles.append(dict(tracteur=tr, dossier="", type="Sans pointage OB",
                                  detail="Aucun pointage ni relevé compteur (affrété) : arrêts considérés réalisés "
                                         "d'après la date planifiée, km uniquement via PTV."))
        affrete = bool(g["affrete"].any())
        societe = next((s for s in g["societe"] if s), "")
        lignes.append(dict(
            tracteur=tr, type_ressource=("Affrété" if affrete else "Tracteur"), societe=societe,
            chauffeurs=", ".join(chauffeurs), plateaux=", ".join(plateaux),
            debut=g["ts_plan"].min(), fin=g["ts_plan"].max(), jours=jours,
            dossiers=g["dossier"].nunique(), arrets=len(g), ca=round(ca, 2),
            km_charge=round(km_tot - km_vide, 1), km_vide=round(km_vide, 1), km_total=round(km_tot, 1),
            pct_vide=round(km_vide / km_tot * 100, 1) if km_tot else np.nan,
            renta=round(ca / km_tot, 3) if km_tot else np.nan,
            ca_jour=round(ca / jours, 0) if jours else np.nan,
            km_groupage=round(km_group, 1), km_compteur=km_cpt, ecart_compteur=ecart,
            sans_ca=int((~dt.loc[dt["tracteur"] == tr, "ca_trouve"]).sum()),
        ))
    return pd.DataFrame(lignes)


# ══════════════════════════════════════════════════════════════════
#  EXPORT EXCEL
# ══════════════════════════════════════════════════════════════════

COLS_FLOTTE = {
    "tracteur": "Tracteur / ressource", "type_ressource": "Type", "societe": "Société affrétée",
    "chauffeurs": "Chauffeur(s)", "plateaux": "Plateau(x)",
    "jours": "Jours", "dossiers": "Dossiers", "ca": "CA (€)", "km_charge": "KM chargés",
    "km_vide": "KM à vide", "km_total": "KM total", "pct_vide": "% vide", "renta": "Renta €/km",
    "ca_jour": "CA / jour (€)", "sans_ca": "Dossiers sans CA", "km_compteur": "KM compteur", "ecart_compteur": "Écart compteur (%)",
}
COLS_DOSSIERS = {
    "tracteur": "Tracteur", "chauffeur": "Chauffeur", "dossier": "N° dossier", "client": "Client",
    "chargement": "Chargement", "dechargement": "Déchargement", "produit": "Produit", "poids": "Poids (t)",
    "ca": "CA (€)", "km_charge": "KM chargés", "km_vide": "KM à vide", "km_total": "KM total",
    "renta": "Renta €/km", "groupage": "Groupage", "realise": "Réalisé", "etat_vente": "État vente",
    "vente_planning": "Vente planning",
}
COLS_TRAJETS = {
    "tracteur": "Tracteur", "ordre": "Ordre", "de_act": "De (act.)", "de_nom": "De", "de_ville": "Ville départ",
    "a_act": "À (act.)", "a_nom": "À", "a_ville": "Ville arrivée", "type": "Type", "bord": "À bord",
    "km": "KM PTV", "km_compteur": "KM compteur", "a_ts": "Arrivée",
}
COLS_CTRL = {"type": "Contrôle", "tracteur": "Tracteur", "dossier": "N° dossier", "detail": "Détail"}


def _traj_affichage(df_traj):
    if df_traj.empty:
        return df_traj
    d = df_traj.copy()
    d["de_ville"] = d["de"].map(lambda l: f"{l[0]} ({l[2]}-{l[1]})")
    d["a_ville"] = d["a"].map(lambda l: f"{l[0]} ({l[2]}-{l[1]})")
    d["bord"] = d["a_bord"].map(lambda t: ", ".join(t) if t else "—")
    return d


def export_excel(df_flotte, df_dos, df_traj, df_ctrl):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        for nom, df, cols in [("Flotte", df_flotte, COLS_FLOTTE), ("Dossiers", df_dos, COLS_DOSSIERS),
                              ("Trajets", _traj_affichage(df_traj), COLS_TRAJETS), ("Contrôles", df_ctrl, COLS_CTRL)]:
            if df is None or df.empty:
                continue
            if nom == "Dossiers" and "premier_arret" in df.columns:
                # Ordre de la tournée : par tracteur, puis par premier arrêt du dossier
                df = df.sort_values(["tracteur", "premier_arret", "dossier"], kind="stable")
            c = [k for k in cols if k in df.columns]
            x = df[c].rename(columns=cols).copy()
            for col in x.columns:
                if pd.api.types.is_datetime64_any_dtype(x[col]):
                    x[col] = x[col].dt.strftime("%d/%m/%Y %H:%M").fillna("")
                if x[col].dtype == bool:
                    x[col] = x[col].map({True: "Oui", False: "Non"})
            x.to_excel(w, sheet_name=nom, index=False)
            _style(w.sheets[nom], len(x))
    return out.getvalue()


def _style(ws, n):
    head = PatternFill("solid", fgColor="003087")
    alt = PatternFill("solid", fgColor="EEF2F7")
    for cell in ws[1]:
        cell.fill, cell.font = head, Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(2, n + 2):
        if r % 2 == 0:
            for cell in ws[r]:
                cell.fill = alt
    for col in ws.columns:
        L = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(L + 3, 9), 55)
    ws.freeze_panes = "A2"


# ══════════════════════════════════════════════════════════════════
#  INTERFACE
# ══════════════════════════════════════════════════════════════════

def _fmt_eur(v):
    return "—" if v is None or pd.isna(v) else f"{v:,.0f} €".replace(",", " ")


def _fmt_km(v):
    return "—" if v is None or pd.isna(v) else f"{v:,.0f}".replace(",", " ")


def _js(v):
    """Valeur → JSON (NaN / NaT → null, Timestamp → 'JJ/MM HH:MM')."""
    if v is None:
        return None
    if isinstance(v, (pd.Timestamp,)):
        return None if pd.isna(v) else v.strftime("%d/%m %H:%M")
    if isinstance(v, (float, np.floating)):
        return None if np.isnan(v) else round(float(v), 3)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, (tuple, list)):
        return [_js(x) for x in v]
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return v


def _records(df, cols):
    if df is None or df.empty:
        return []
    return [{c: _js(r.get(c)) for c in cols} for r in df.to_dict("records")]


def construire_tableau_de_bord(df_flotte, df_dos, df_traj, df_ctrl, stops, meta: dict) -> str:
    """Page HTML autonome (style hub) : vue d'ensemble, camion, dossiers, contrôles."""
    trajets = {}
    if not df_traj.empty:
        for r in df_traj.sort_values(["tracteur", "ordre"]).to_dict("records"):
            trajets.setdefault(r["tracteur"], []).append(dict(
                o=_js(r["ordre"]), t=r["type"], km=_js(r["km"]), cpt=_js(r["km_compteur"]),
                dn=r["de_nom"], dv=r["de"][0], dp=r["de"][2], da=r["de_act"],
                an=r["a_nom"], av=r["a"][0], ap=r["a"][2], aa=r["a_act"],
                bord=list(r["a_bord"]), arr=_js(r["a_ts"]),
                la=_js(r["lat_de"]), lo=_js(r["lon_de"]), lb=_js(r["lat_a"]), lob=_js(r["lon_a"]),
                grp=bool(r.get("releve_groupe", False))))
    arrets = {}
    for r in stops.sort_values(["tracteur", "ordre"]).to_dict("records"):
        site = None
        if pd.notna(r["ts_arr"]) and pd.notna(r["ts_dep"]):
            site = round((r["ts_dep"] - r["ts_arr"]).total_seconds() / 3600, 1)
        arrets.setdefault(r["tracteur"], []).append(dict(
            o=_js(r["ordre"]), d=r["dossier"], act=r["activite"], nom=r["nom"], v=r["ville"], p=r["pays"],
            cp=r["cp"], prevu=_js(r["ts_plan"]), arr=_js(r["ts_arr"]), dep=_js(r["ts_dep"]), site=site,
            odo=_js(r["odo_km"]), ch=r["chauffeur"], pl=r["plateau"], real=bool(r["realise"])))
    payload = dict(
        meta=meta,
        flotte=_records(df_flotte, list(COLS_FLOTTE) + ["km_groupage", "debut", "fin", "arrets"]),
        dossiers=_records(df_dos, list(COLS_DOSSIERS) + ["ca_trouve", "date_debut"]),
        controles=_records(df_ctrl, list(COLS_CTRL)),
        trajets=trajets, arrets=arrets,
        seuils=dict(vide=SEUIL_ALERTE_VIDE, site=SEUIL_TEMPS_SITE, odo=SEUIL_ECART_ODO),
    )
    data = json.dumps(payload, ensure_ascii=False, default=str).replace("</", "<\\/")
    a = _assets_dashboard()
    return (DASH_HTML.replace("__FONTS__", a["fonts"], 1).replace("__LEAFLET_CSS__", a["css"], 1)
            .replace("__LEAFLET_JS__", a["js"], 1).replace("__DATA__", data, 1))


ASSETS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "renta_plateau")


@st.cache_resource(show_spinner=False)
def _assets_dashboard():
    """Leaflet et la police Inter embarqués dans la page : aucune dépendance à un CDN (iframe, réseau d'entreprise)."""
    import base64
    lire = lambda n: open(os.path.join(ASSETS_DIR, n), encoding="utf-8").read()
    fonts = ""
    for w in (400, 500, 600, 700):
        b64 = base64.b64encode(open(os.path.join(ASSETS_DIR, f"inter-latin-{w}-normal.woff2"), "rb").read()).decode()
        fonts += (f"@font-face{{font-family:'Inter';font-style:normal;font-weight:{w};font-display:swap;"
                  f"src:url(data:font/woff2;base64,{b64}) format('woff2')}}")
    return {"fonts": fonts, "css": lire("leaflet.css"), "js": lire("leaflet.js").replace("</script", "<\\/script")}


def bouton_plein_ecran(page_html: str):
    contenu = json.dumps(page_html).replace("</", "<\\/")
    components.html(f"""
<style>
  body {{ margin: 0; font-family: -apple-system, BlinkMacSystemFont, "SF Pro Text", "Inter", "Helvetica Neue", Arial, sans-serif; }}
  button {{ height: 42px; padding: 0 22px; border: none; border-radius: 980px; background: #e8e8ed; color: #0071e3;
           font: 500 15px -apple-system, BlinkMacSystemFont, "SF Pro Text", "Inter", "Helvetica Neue", Arial, sans-serif;
           cursor: pointer; }}
  button:hover {{ background: #dedee3; }}
  span {{ margin-left: 10px; font-size: 13px; color: #6e6e73; }}
</style>
<button id="b">Ouvrir en plein écran</button><span id="m"></span>
<script>
const PAGE = {contenu};
document.getElementById('b').onclick = () => {{
  const w = window.open(URL.createObjectURL(new Blob([PAGE], {{type: 'text/html'}})), '_blank');
  if (!w) document.getElementById('m').textContent = 'Pop-up bloquée : autorisez-les pour le HUB.';
}};
</script>""", height=50)


DASH_HTML = r"""<!doctype html>
<html lang="fr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Rentabilité plateau</title>
<style>__FONTS__</style>
<style>__LEAFLET_CSS__</style>
<script>__LEAFLET_JS__</script>
<style>
:root{--bg:#f5f5f7;--card:#fff;--text:#1d1d1f;--sub:#6e6e73;--line:#d2d2d7;--hair:#e5e5ea;--soft:#e8e8ed;
--blue:#0071e3;--green:#248a3d;--green-bg:#e3f5e8;--orange:#c93400;--orange-bg:#fff1e6;--red:#d70015;--red-bg:#fde8ea;
--c-ch:#0071e3;--c-vide:#a1a1a6;--c-grp:#ff9500;--c-rec:#af52de;
--shadow:0 1px 2px rgba(0,0,0,.04),0 0 0 .5px rgba(0,0,0,.06);
--font:-apple-system,BlinkMacSystemFont,"SF Pro Text","SF Pro Display","Inter","Helvetica Neue",Arial,sans-serif}
body,button,input,select,table{font-family:var(--font)}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--text);font-family:var(--font);-webkit-font-smoothing:antialiased;font-size:14px}
.wrap{max-width:1180px;margin:0 auto;padding:22px 20px 40px}
.head{display:flex;align-items:flex-end;justify-content:space-between;gap:16px;flex-wrap:wrap;margin-bottom:18px}
.head h1{font-size:28px;font-weight:700;letter-spacing:-.02em;margin:0}
.head .meta{color:var(--sub);font-size:13px;margin-top:4px}
.seg{display:inline-flex;background:var(--soft);border-radius:10px;padding:2px}
.seg button{border:0;background:transparent;font:500 13px var(--font);color:var(--text);padding:7px 14px;border-radius:8px;cursor:pointer}
.seg button.on{background:#fff;box-shadow:0 1px 3px rgba(0,0,0,.12)}
.kpis{display:grid;grid-template-columns:repeat(6,1fr);gap:12px;margin-bottom:18px}
.kpi{background:var(--card);border-radius:16px;padding:14px 16px;box-shadow:var(--shadow)}
.kpi .l{font-size:12px;color:var(--sub);font-weight:500}
.kpi .v{font-size:24px;font-weight:600;letter-spacing:-.02em;margin-top:4px;white-space:nowrap}
.kpi .s{font-size:12px;color:var(--sub);margin-top:2px}
.card{background:var(--card);border-radius:18px;box-shadow:var(--shadow);padding:18px 20px;margin-bottom:16px}
.card h2{font-size:17px;font-weight:600;margin:0 0 4px;letter-spacing:-.01em}
.card .hint{font-size:12.5px;color:var(--sub);margin:0 0 14px}
.row-head{display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap;margin-bottom:12px}
.legend{display:flex;gap:14px;font-size:12px;color:var(--sub);flex-wrap:wrap}
.legend i{display:inline-block;width:10px;height:10px;border-radius:3px;margin-right:5px;vertical-align:-1px}
/* Classement */
.rank{display:grid;grid-template-columns:28px minmax(150px,1.3fr) 92px minmax(180px,2fr) 96px 64px 90px;align-items:center;
  gap:12px;padding:11px 8px;border-top:.5px solid var(--hair);cursor:pointer;border-radius:10px}
.rank:hover{background:#f7f7f9}
.rank.h{cursor:default;border-top:0;font-size:11.5px;color:var(--sub);font-weight:500;padding-top:0}
.rank.h:hover{background:none}
.rank .n{color:var(--sub);font-variant-numeric:tabular-nums;text-align:right}
.rank .tr b{display:block;font-weight:600}.rank .tr span{font-size:12px;color:var(--sub)}
.renta{font-size:18px;font-weight:600;font-variant-numeric:tabular-nums;letter-spacing:-.01em}
.up{color:var(--green)}.down{color:var(--orange)}
.bar{display:flex;height:10px;border-radius:5px;overflow:hidden;background:var(--soft)}
.bar i{display:block;height:100%}
.barlbl{font-size:11.5px;color:var(--sub);margin-top:4px;font-variant-numeric:tabular-nums}
.num{text-align:right;font-variant-numeric:tabular-nums}
.badges{display:flex;gap:4px;justify-content:flex-end;flex-wrap:wrap}
.b{display:inline-block;font-size:11px;font-weight:500;padding:2px 7px;border-radius:6px;white-space:nowrap}
.b.o{background:var(--orange-bg);color:var(--orange)}.b.r{background:var(--red-bg);color:var(--red)}
.b.g{background:var(--green-bg);color:var(--green)}.b.n{background:var(--soft);color:var(--sub)}
.b.p{background:#f4ebfb;color:#8944ab}
/* Clients */
.cli{display:grid;grid-template-columns:minmax(140px,1fr) 2fr 90px 70px;gap:12px;align-items:center;padding:8px 0;border-top:.5px solid var(--hair)}
.cli:first-of-type{border-top:0}
/* Camion */
.trk-top{display:flex;gap:10px;align-items:center;flex-wrap:wrap;margin-bottom:14px}
select,input[type=search]{font:14px var(--font);padding:8px 12px;border-radius:10px;border:1px solid var(--line);background:#fff;color:var(--text)}
.nav{border:0;background:var(--soft);color:var(--blue);width:34px;height:34px;border-radius:50%;font-size:16px;cursor:pointer}
.nav:hover{background:#dedee3}
.split{display:grid;grid-template-columns:1.15fr 1fr;gap:16px}
#map{height:560px;border-radius:14px;overflow:hidden;background:#eef0f3}
.tl{max-height:560px;overflow:auto;padding-right:6px}
.stop{display:grid;grid-template-columns:46px 1fr auto;gap:10px;align-items:start;padding:6px 0}
.stop .act{font-size:11px;font-weight:700;color:#fff;border-radius:7px;text-align:center;padding:4px 0;margin-top:2px}
.act.CH{background:var(--green)}.act.DECH{background:var(--blue)}
.stop .pl b{font-weight:600}.stop .pl div{font-size:12px;color:var(--sub);margin-top:1px}
.stop .tm{text-align:right;font-size:12px;color:var(--sub);white-space:nowrap;font-variant-numeric:tabular-nums}
.stop.prev{opacity:.55}
.leg{display:grid;grid-template-columns:46px 1fr;gap:10px;align-items:center;min-height:30px}
.leg .ln{justify-self:center;width:3px;height:100%;min-height:30px;border-radius:2px}
.leg .lt{font-size:12px;color:var(--sub);font-variant-numeric:tabular-nums}
.leg .lt b{color:var(--text);font-weight:600}
/* Tables */
table{width:100%;border-collapse:collapse;font-size:13px}
th{text-align:left;font-weight:500;color:var(--sub);font-size:12px;padding:8px 8px;border-bottom:.5px solid var(--line);
  cursor:pointer;white-space:nowrap;user-select:none;position:sticky;top:0;background:#fff}
th.r,td.r{text-align:right}
td{padding:9px 8px;border-bottom:.5px solid var(--hair);font-variant-numeric:tabular-nums;vertical-align:top}
tr:hover td{background:#fafafc}
.tablewrap{max-height:620px;overflow:auto}
.filters{display:flex;gap:10px;flex-wrap:wrap;align-items:center}
.ctrl-grp{margin-bottom:14px}
.ctrl-grp h3{font-size:14px;font-weight:600;margin:0 0 6px;display:flex;align-items:center;gap:8px}
.ctrl-item{display:grid;grid-template-columns:110px 100px 1fr;gap:10px;font-size:13px;padding:7px 0;border-top:.5px solid var(--hair)}
.ctrl-item span{color:var(--sub)}
.empty{color:var(--sub);padding:20px 0;text-align:center}
.hide{display:none}
@media(max-width:900px){.kpis{grid-template-columns:repeat(3,1fr)}.split{grid-template-columns:1fr}
 .rank{grid-template-columns:1fr 80px 90px}.rank>:nth-child(1),.rank>:nth-child(4),.rank>:nth-child(6),.rank>:nth-child(7){display:none}}
</style></head><body><div class="wrap">
<div class="head"><div><h1>Rentabilité plateau</h1><div class="meta" id="meta"></div></div>
<div class="seg" id="tabs"><button data-t="vue" class="on">Vue d'ensemble</button><button data-t="cam">Camion</button>
<button data-t="dos">Dossiers</button><button data-t="ctl">Contrôles</button></div></div>

<section id="t-vue">
 <div class="kpis" id="kpis"></div>
 <div class="card">
  <div class="row-head"><div><h2>Classement des camions</h2>
   <p class="hint" style="margin:0">Renta comparée à la moyenne de la flotte · cliquez sur un camion pour voir sa tournée.</p></div>
   <div class="seg" id="sort"><button data-s="renta" class="on">Renta</button><button data-s="ca">CA</button>
   <button data-s="km_total">KM</button><button data-s="pct_vide">% vide</button></div></div>
  <div class="legend" style="margin-bottom:6px"><span><i style="background:var(--c-ch)"></i>Chargé</span>
   <span><i style="background:var(--c-grp)"></i>Groupage</span><span><i style="background:var(--c-vide)"></i>À vide</span></div>
  <div id="rank"></div>
 </div>
 <div class="card"><h2>Par client</h2><p class="hint">CA et renta des dossiers de la période, km imputés compris · les dossiers sans CA sont exclus de la renta.</p><div id="clients"></div></div>
</section>

<section id="t-cam" class="hide">
 <div class="trk-top"><button class="nav" id="prv">‹</button><select id="trsel"></select><button class="nav" id="nxt">›</button>
  <span id="trinfo" style="color:var(--sub);font-size:13px"></span></div>
 <div class="kpis" id="tkpis"></div>
 <div class="split">
  <div class="card" style="padding:12px"><div id="map"></div>
   <div class="legend" style="margin:10px 4px 0"><span><i style="background:var(--c-ch)"></i>Chargé</span>
   <span><i style="background:var(--c-grp)"></i>Groupage</span><span><i style="background:var(--c-vide)"></i>À vide</span>
   <span><i style="background:var(--c-rec)"></i>Reconstitué</span><span>Tracés à vol d'oiseau · km PTV</span></div></div>
  <div class="card"><h2>Tournée</h2><p class="hint">Arrêts dans l'ordre, trajets entre deux arrêts, temps sur site.</p><div class="tl" id="tl"></div></div>
 </div>
 <div class="card"><h2>Dossiers du camion</h2><div class="tablewrap" id="tdos"></div></div>
</section>

<section id="t-dos" class="hide">
 <div class="card"><div class="row-head"><h2>Dossiers</h2><div class="filters">
  <input type="search" id="q" placeholder="Dossier, client, ville…"><select id="fcli"><option value="">Tous les clients</option></select>
  <select id="ftr"><option value="">Tous les camions</option></select></div></div>
  <div class="tablewrap" id="dos"></div><div class="meta" id="dosfoot" style="margin-top:10px;color:var(--sub);font-size:12.5px"></div></div>
</section>

<section id="t-ctl" class="hide"><div class="card"><h2>Contrôles</h2><p class="hint">Points à vérifier avant d'exploiter les chiffres.</p><div id="ctl"></div></div></section>
</div>
<script>
const D = __DATA__;
const nf0 = new Intl.NumberFormat('fr-FR',{maximumFractionDigits:0}), nf1 = new Intl.NumberFormat('fr-FR',{minimumFractionDigits:1,maximumFractionDigits:1}),
      nf3 = new Intl.NumberFormat('fr-FR',{minimumFractionDigits:3,maximumFractionDigits:3});
const eur = v => v==null?'—':nf0.format(v)+' €', km = v => v==null?'—':nf0.format(v), r3 = v => v==null?'—':nf3.format(v),
      pct = v => v==null?'—':nf1.format(v)+' %', esc = s => String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const F = D.flotte, sum = (a,k) => a.reduce((s,x)=>s+(x[k]||0),0);
const CA = sum(F,'ca'), KM = sum(F,'km_total'), KV = sum(F,'km_vide'), MOY = KM? CA/KM : null;
const cmp = v => (v==null||MOY==null||F.length<2||Math.abs(v-MOY)<0.0005)?'':(v>MOY?'up':'down');
const COL = {'Chargé':'#0071e3','À vide':'#a1a1a6','Groupage':'#ff9500','Reconstitué':'#af52de'};

document.getElementById('meta').textContent = [D.meta.periode, D.meta.convention, 'calculé le '+D.meta.genere].filter(Boolean).join(' · ');

function kpiHTML(items){ return items.map(([l,v,s])=>`<div class="kpi"><div class="l">${l}</div><div class="v">${v}</div>${s?`<div class="s">${s}</div>`:''}</div>`).join(''); }
const jours = sum(F,'jours'), sansCA = sum(F,'sans_ca'), ndos = D.dossiers.length;
document.getElementById('kpis').innerHTML = kpiHTML([
 ['Chiffre d\'affaires', eur(CA), `${F.length} camion${F.length>1?'s':''}`],
 ['Renta', r3(MOY), '€ par km parcouru'],
 ['Kilomètres', km(KM), `${km(KM-KV)} chargés`],
 ['À vide', pct(KM?KV/KM*100:null), `${km(KV)} km`],
 ['CA par jour', eur(jours?CA/jours:null), `${jours} jours camion`],
 ['Dossiers', nf0.format(ndos), sansCA?`<span style="color:var(--orange)">${sansCA} sans CA</span>`:'tous avec CA']]);

// ── Classement
let sortKey='renta';
function drawRank(){
 const asc = sortKey==='pct_vide', arr=[...F].sort((a,b)=>((asc?1:-1)*((a[sortKey]??-1e9)-(b[sortKey]??-1e9))));
 const maxKm = Math.max(...F.map(f=>f.km_total||0),1);
 let h = `<div class="rank h"><div></div><div>Camion</div><div class="num">Renta €/km</div><div>Kilomètres</div><div class="num">CA</div><div class="num">Jours</div><div></div></div>`;
 arr.forEach((f,i)=>{
  const ch=(f.km_charge||0)-(f.km_groupage||0), g=f.km_groupage||0, v=f.km_vide||0, w=x=>(x/maxKm*100).toFixed(2)+'%';
  const cls = cmp(f.renta);
  const bd=[]; if(f.sans_ca) bd.push(`<span class="b o">${f.sans_ca} sans CA</span>`);
  if(f.ecart_compteur!=null && Math.abs(f.ecart_compteur)>D.seuils.odo) bd.push(`<span class="b r">cpt ${f.ecart_compteur>0?'+':''}${nf1.format(f.ecart_compteur)} %</span>`);
  if((f.plateaux||'').includes(',')) bd.push(`<span class="b p">accroche</span>`);
  if(f.type_ressource==='Affrété') bd.push(`<span class="b n">${esc(f.societe||'affrété')}</span>`);
  h+=`<div class="rank" data-tr="${esc(f.tracteur)}"><div class="n">${i+1}</div>
   <div class="tr"><b>${esc(f.tracteur)}</b><span>${esc(f.chauffeurs||'—')}</span></div>
   <div class="num renta ${cls}">${r3(f.renta)}</div>
   <div><div class="bar"><i style="width:${w(ch)};background:var(--c-ch)"></i><i style="width:${w(g)};background:var(--c-grp)"></i><i style="width:${w(v)};background:var(--c-vide)"></i></div>
    <div class="barlbl">${km(f.km_total)} km · ${pct(f.pct_vide)} vide</div></div>
   <div class="num">${eur(f.ca)}</div><div class="num">${f.jours??'—'}</div><div class="badges">${bd.join('')}</div></div>`;
 });
 const el=document.getElementById('rank'); el.innerHTML=h;
 el.querySelectorAll('.rank[data-tr]').forEach(r=>r.onclick=()=>{ showTab('cam'); selTruck(r.dataset.tr); });
}
document.querySelectorAll('#sort button').forEach(b=>b.onclick=()=>{document.querySelectorAll('#sort button').forEach(x=>x.classList.remove('on'));b.classList.add('on');sortKey=b.dataset.s;drawRank();});
drawRank();

// ── Clients
(function(){
 const m={}; D.dossiers.forEach(d=>{const c=d.client||'—'; m[c]=m[c]||{ca:0,km:0,kmca:0,n:0,sans:0}; m[c].n++; m[c].km+=d.km_total||0;
  if(d.ca_trouve===false) m[c].sans++; else {m[c].ca+=d.ca||0; m[c].kmca+=d.km_total||0;}});
 const arr=Object.entries(m).sort((a,b)=>b[1].ca-a[1].ca), mx=Math.max(...arr.map(x=>x[1].ca),1);
 document.getElementById('clients').innerHTML = arr.length? arr.map(([c,v])=>{const r=v.kmca?v.ca/v.kmca:null, cls=cmp(r);
  return `<div class="cli"><div><b style="font-weight:600">${esc(c)}</b><div style="font-size:12px;color:var(--sub)">${v.n} dossier${v.n>1?'s':''} · ${km(v.km)} km${v.sans?` · <span style="color:var(--orange)">${v.sans} sans CA</span>`:''}</div></div>
  <div class="bar" style="height:8px"><i style="width:${(v.ca/mx*100).toFixed(1)}%;background:var(--blue)"></i></div>
  <div class="num">${eur(v.ca)}</div><div class="num ${cls}" style="font-weight:600">${r3(r)}</div></div>`}).join('') : '<div class="empty">Aucun dossier.</div>';
})();

// ── Tableaux triables
function table(el, cols, rows){
 let key=null, dir=-1;
 function draw(){
  const rs = key? [...rows].sort((a,b)=>{const x=a[key],y=b[key]; if(x==null) return 1; if(y==null) return -1; return (typeof x==='number'?x-y:String(x).localeCompare(String(y),'fr'))*dir;}) : rows;
  el.innerHTML = rs.length? `<table><thead><tr>${cols.map(c=>`<th class="${c.r?'r':''}" data-k="${c.k}">${c.l}${key===c.k?(dir>0?' ↑':' ↓'):''}</th>`).join('')}</tr></thead>
   <tbody>${rs.map(r=>`<tr>${cols.map(c=>`<td class="${c.r?'r':''}">${c.f?c.f(r[c.k],r):esc(r[c.k]??'—')}</td>`).join('')}</tr>`).join('')}</tbody></table>` : '<div class="empty">Aucun dossier.</div>';
  el.querySelectorAll('th').forEach(th=>th.onclick=()=>{ if(key===th.dataset.k) dir=-dir; else {key=th.dataset.k; dir=-1;} draw(); });
 }
 draw();
}
const rentaCell = v => v==null?'<span style="color:var(--sub)">—</span>':`<b class="${cmp(v)}" style="font-weight:600">${r3(v)}</b>`;
const dosBadges = (v,r) => [r.groupage?'<span class="b o">groupage</span>':'', r.ca_trouve===false?'<span class="b r">sans CA</span>':'', r.realise===false?'<span class="b n">prévu</span>':''].join(' ');
const COLS_DOS = [
 {k:'dossier',l:'Dossier'},{k:'tracteur',l:'Camion'},{k:'client',l:'Client'},
 {k:'chargement',l:'Chargement'},{k:'dechargement',l:'Déchargement'},
 {k:'ca',l:'CA',r:1,f:eur},{k:'km_charge',l:'KM ch.',r:1,f:km},{k:'km_vide',l:'KM vide',r:1,f:km},
 {k:'renta',l:'Renta',r:1,f:rentaCell},{k:'groupage',l:'',f:dosBadges}];

// ── Dossiers
(function(){
 const cl=[...new Set(D.dossiers.map(d=>d.client).filter(Boolean))].sort(), tr=F.map(f=>f.tracteur);
 document.getElementById('fcli').innerHTML += cl.map(c=>`<option>${esc(c)}</option>`).join('');
 document.getElementById('ftr').innerHTML += tr.map(c=>`<option>${esc(c)}</option>`).join('');
 const go=()=>{ const q=document.getElementById('q').value.toLowerCase(), c=document.getElementById('fcli').value, t=document.getElementById('ftr').value;
  const rows=D.dossiers.filter(d=>(!c||d.client===c)&&(!t||d.tracteur===t)&&(!q||[d.dossier,d.client,d.chargement,d.dechargement,d.chauffeur].join(' ').toLowerCase().includes(q)));
  table(document.getElementById('dos'), COLS_DOS, rows);
  const ca=sum(rows,'ca'), k=sum(rows,'km_total');
  document.getElementById('dosfoot').textContent = `${rows.length} dossier(s) · CA ${eur(ca)} · ${km(k)} km · renta ${r3(k?ca/k:null)} €/km`; };
 ['q','fcli','ftr'].forEach(id=>document.getElementById(id).addEventListener('input',go)); go();
})();

// ── Contrôles
(function(){
 const g={}; D.controles.forEach(c=>{(g[c.type]=g[c.type]||[]).push(c);});
 const sev = t => /Vide|Écart|échoué|sans CA/.test(t)?'r':(/Accroche/.test(t)?'p':'o');
 document.getElementById('ctl').innerHTML = Object.keys(g).length? Object.entries(g).sort((a,b)=>b[1].length-a[1].length).map(([t,l])=>
  `<div class="ctrl-grp"><h3><span class="b ${sev(t)}">${l.length}</span>${esc(t)}</h3>${l.map(c=>`<div class="ctrl-item"><span>${esc(c.tracteur||'—')}</span><span>${esc(c.dossier||'')}</span><div>${esc(c.detail)}</div></div>`).join('')}</div>`).join('')
  : '<div class="empty">Aucune anomalie détectée.</div>';
 const btn=document.querySelector('[data-t=ctl]'); if(D.controles.length) btn.textContent=`Contrôles (${D.controles.length})`;
})();

// ── Camion
let map=null, layer=null, cur=null;
const sel=document.getElementById('trsel');
sel.innerHTML = F.map(f=>`<option value="${esc(f.tracteur)}">${esc(f.tracteur)} — ${esc(f.chauffeurs||'')}</option>`).join('');
sel.onchange=()=>selTruck(sel.value);
document.getElementById('prv').onclick=()=>{const i=F.findIndex(f=>f.tracteur===cur); selTruck(F[(i-1+F.length)%F.length].tracteur);};
document.getElementById('nxt').onclick=()=>{const i=F.findIndex(f=>f.tracteur===cur); selTruck(F[(i+1)%F.length].tracteur);};
function selTruck(t){
 cur=t; sel.value=t; const f=F.find(x=>x.tracteur===t); if(!f) return;
 document.getElementById('trinfo').textContent = `Plateau ${f.plateaux||'—'} · ${f.dossiers} dossiers · ${f.arrets} arrêts`;
 const cls = cmp(f.renta);
 document.getElementById('tkpis').innerHTML = kpiHTML([
  ['CA', eur(f.ca), `${eur(f.ca_jour)} / jour`],
  ['Renta', `<span class="${cls}">${r3(f.renta)}</span>`, MOY?`moyenne flotte ${r3(MOY)}`:''],
  ['KM chargés', km(f.km_charge), f.km_groupage?`dont ${km(f.km_groupage)} en groupage`:''],
  ['KM à vide', km(f.km_vide), pct(f.pct_vide)],
  ['Compteur', km(f.km_compteur), f.ecart_compteur==null?'':`${f.ecart_compteur>0?'+':''}${nf1.format(f.ecart_compteur)} % vs PTV`],
  ['Jours', f.jours??'—', f.sans_ca?`<span style="color:var(--orange)">${f.sans_ca} dossier(s) sans CA</span>`:'']]);
 drawTimeline(t); drawMap(t);
 table(document.getElementById('tdos'), COLS_DOS.filter(c=>c.k!=='tracteur'), D.dossiers.filter(d=>d.tracteur===t));
}
function drawTimeline(t){
 const A=D.arrets[t]||[], T=D.trajets[t]||[], byO={}; T.forEach(x=>byO[x.o]=x);
 const legHTML = x => { if(!x) return ''; const c=COL[x.t]||'#999';
  const extra = x.t==='À vide'&&x.km>D.seuils.vide?' <span class="b r">à vérifier</span>':'';
  return `<div class="leg"><div class="ln" style="background:${c}"></div><div class="lt"><b>${km(x.km)} km</b> · ${x.t.toLowerCase()}${x.bord.length?' · '+x.bord.join(', '):''}${x.cpt!=null?(x.grp?' · <span title="Relevés compteur pris au même moment : valeur non exploitable">compteur groupé</span>':` · compteur ${km(x.cpt)}`):''}${extra}</div></div>`; };
 let h='';
 const pre=byO[0]; if(pre) h+=`<div class="stop prev"><div class="act CH">CH*</div><div class="pl"><b>${esc(pre.dn||pre.dv)}</b><div>${esc(pre.dv)} (${esc(pre.dp)}) · chargement avant la période</div></div><div class="tm"></div></div>`+legHTML(pre);
 A.forEach((a,i)=>{
  const site = a.site!=null?`<div>${nf1.format(a.site)} h sur site</div>`:'';
  const warn = a.site!=null&&a.site>D.seuils.site?' style="color:var(--orange)"':'';
  h+=`<div class="stop ${a.real?'':'prev'}"><div class="act ${a.act}">${a.act}</div>
   <div class="pl"><b>${esc(a.nom)}</b><div>${esc(a.v)} (${esc(a.p)}-${esc(a.cp)}) · ${esc(a.d)}${a.real?'':' · prévu'}</div></div>
   <div class="tm"${warn}><div>${a.arr||a.prevu||''}</div>${site}</div></div>`;
  h+=legHTML(byO[a.o]);
 });
 document.getElementById('tl').innerHTML = h || '<div class="empty">Aucun arrêt.</div>';
}
function drawMap(t){
 if(typeof L==='undefined'){ document.getElementById('map').innerHTML='<div class="empty">Carte indisponible (hors ligne).</div>'; return; }
 if(!map){ map=L.map('map',{zoomControl:true,attributionControl:true});
  const esri = s => `https://server.arcgisonline.com/ArcGIS/rest/services/${s}/MapServer/tile/{z}/{y}/{x}`;
  L.tileLayer(esri('Canvas/World_Light_Gray_Base'),{attribution:'Tiles © Esri',maxNativeZoom:16,maxZoom:18}).addTo(map);
  L.tileLayer(esri('Canvas/World_Light_Gray_Reference'),{maxNativeZoom:16,maxZoom:18}).addTo(map); }
 if(layer) layer.remove(); layer=L.featureGroup().addTo(map);
 const T=(D.trajets[t]||[]).filter(x=>x.la!=null&&x.lb!=null), pts={};
 T.forEach(x=>{
  L.polyline([[x.la,x.lo],[x.lb,x.lob]],{color:COL[x.t]||'#999',weight:x.t==='À vide'?3:4,opacity:.9,dashArray:x.t==='À vide'?'6 6':null})
   .bindTooltip(`${esc(x.dv)} → ${esc(x.av)}<br><b>${km(x.km)} km</b> · ${x.t}${x.bord.length?'<br>'+x.bord.join(', '):''}`,{sticky:true}).addTo(layer);
  pts[x.la+','+x.lo]=[x.la,x.lo,x.dn||x.dv]; pts[x.lb+','+x.lob]=[x.lb,x.lob,x.an||x.av];
 });
 Object.values(pts).forEach(([a,b,n])=>L.circleMarker([a,b],{radius:5,color:'#fff',weight:2,fillColor:'#1d1d1f',fillOpacity:1}).bindTooltip(esc(n)).addTo(layer));
 setTimeout(()=>{ map.invalidateSize(); if(layer.getLayers().length) map.fitBounds(layer.getBounds(),{padding:[30,30]}); else map.setView([48.5,4],5); },50);
}

// ── Onglets
function showTab(t){ document.querySelectorAll('#tabs button').forEach(b=>b.classList.toggle('on',b.dataset.t===t));
 ['vue','cam','dos','ctl'].forEach(x=>document.getElementById('t-'+x).classList.toggle('hide',x!==t));
 if(t==='cam'){ if(!cur&&F.length) selTruck(F[0].tracteur); else if(map) setTimeout(()=>map.invalidateSize(),50); } }
document.querySelectorAll('#tabs button').forEach(b=>b.onclick=()=>showTab(b.dataset.t));
</script></body></html>"""


def main():
    from core import ui

    ui.page("renta_plateau", "Exports planning des plateaux : tournées recomposées par tracteur, "
                             "CA par dossier, km chargés et à vide, rentabilité au kilomètre.")

    if not PTV_API_KEY or PTV_API_KEY == "METS_TA_CLE_ICI":
        st.error("Clé PTV_API_KEY non configurée : le calcul des distances ne fonctionnera pas.")

    if st.session_state.get("rp_schema") != RESULT_SCHEMA:
        for k in [k for k in st.session_state if k.startswith("rp_")]:
            st.session_state.pop(k, None)
        st.session_state["rp_schema"] = RESULT_SCHEMA

    # ── Fichiers ──
    ui.section("Fichiers")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("#### Exports planning")
        st.caption("Un ou plusieurs fichiers Traject_xxxxx.xlsx. Les lignes d'un même tracteur sont "
                   "recomposées d'un fichier à l'autre ; les doublons sont retirés.")
        files_pl = st.file_uploader("Plannings (.xlsx)", type=["xlsx"], accept_multiple_files=True,
                                    key="rp_up_pl", label_visibility="collapsed")
    with c2:
        st.markdown("#### Fichier CA")
        st.caption("N° Dossier · Total des ventes · Client facturation · Etat vente")
        file_ca = st.file_uploader("CA (.xlsx)", type=["xlsx"], key="rp_up_ca", label_visibility="collapsed")

    if not files_pl:
        st.markdown(
            '<ul class="ap-list">'
            '<li><b>1. Plannings</b> <span>— exports Traject, un par tracteur ou chauffeur affrété (AFF-xxx), ou plusieurs pour le même.</span></li>'
            '<li><b>2. CA</b> <span>— export habituel, jointure sur le N° dossier.</span></li>'
            '<li><b>3. Contrôle</b> <span>— tracteurs détectés, chauffeurs, plateaux, dossiers incomplets.</span></li>'
            '<li><b>4. Calcul PTV</b> <span>— km par trajet, imputation chargé / vide / groupage.</span></li>'
            '<li><b>5. Résultats</b> <span>— flotte, tournée par camion, dossiers, contrôles, export Excel.</span></li>'
            '</ul>', unsafe_allow_html=True)
        return

    stops, erreurs, doublons = charger_plannings(files_pl)
    for e in erreurs:
        st.error(f"Lecture impossible — {e}")
    if stops.empty:
        return
    ca_dict = {}
    if file_ca:
        try:
            ca_dict = parse_ca(file_ca)
        except Exception as e:
            st.error(f"Fichier CA : {e}")

    # ── Périmètre ──
    ui.section("Périmètre")
    p1, p2, p3 = st.columns([1.2, 1, 1])
    with p1:
        dmin, dmax = stops["ts_plan"].min(), stops["ts_plan"].max()
        if pd.notna(dmin) and pd.notna(dmax):
            periode = st.date_input("Période (date planifiée de l'arrêt)", value=(dmin.date(), dmax.date()),
                                    min_value=dmin.date(), max_value=dmax.date(), format="DD/MM/YYYY")
        else:
            periode = None
    with p2:
        types = sorted(t for t in stops["type_trans"].unique() if t)
        f_types = st.multiselect("Type de transport", types, default=[t for t in types if t == "PLA"] or types)
    with p3:
        tracteurs = sorted(t for t in stops["tracteur"].unique() if t)
        f_tr = st.multiselect("Tracteurs", tracteurs, default=[], placeholder="Tous")

    o1, o2, o3 = st.columns([1.2, 1, 1])
    with o1:
        convention = st.radio("Km à vide imputés au", list(CONVENTIONS_VIDE), format_func=CONVENTIONS_VIDE.get,
                              horizontal=False)
    with o2:
        realise_seul = st.toggle("Arrêts réalisés uniquement", value=True,
                                 help="Exclut les arrêts planifiés sans pointage OB (trajets à venir).")
    with o3:
        reconstituer = st.toggle("Reconstituer les dossiers coupés", value=True,
                                 help="Ajoute le trajet chargé manquant depuis Origine (dossier commencé avant "
                                      "la période) ou vers Destination (dossier terminé après).")

    sel = stops.copy()
    if periode and isinstance(periode, (tuple, list)) and len(periode) == 2:
        a, b = pd.Timestamp(periode[0]), pd.Timestamp(periode[1]) + pd.Timedelta(days=1)
        sel = sel[(sel["ts_plan"] >= a) & (sel["ts_plan"] < b) | sel["ts_plan"].isna()]
    if f_types:
        sel = sel[sel["type_trans"].isin(f_types) | (sel["type_trans"] == "")]
    if f_tr:
        sel = sel[sel["tracteur"].isin(f_tr)]
    if realise_seul:
        sel = sel[sel["realise"]]
    sel = sel[sel["tracteur"] != ""]

    nb_dos = sel["dossier"].nunique()
    nb_ca = sum(1 for d in sel["dossier"].unique() if d in ca_dict)
    k = st.columns(5)
    k[0].metric("Tracteurs", sel["tracteur"].nunique())
    k[1].metric("Dossiers", nb_dos)
    k[2].metric("Arrêts", len(sel))
    k[3].metric("Dossiers avec CA", f"{nb_ca}/{nb_dos}" if file_ca else "—")
    k[4].metric("Doublons retirés", doublons)
    if not file_ca:
        st.info("Ajoutez le fichier CA pour obtenir la rentabilité. Sans lui, seuls les km sont calculés.")
    if sel.empty:
        st.warning("Aucun arrêt dans le périmètre choisi.")
        return

    with st.expander("Tracteurs et affrétés détectés"):
        det = sel.sort_values("ordre").groupby("tracteur").agg(
            type_ressource=("affrete", lambda s: "Affrété" if s.any() else "Tracteur"),
            societe=("societe", lambda s: next((x for x in s if x), "")),
            chauffeurs=("chauffeur", lambda s: ", ".join(dict.fromkeys(x for x in s if x))),
            plateaux=("plateau", lambda s: ", ".join(dict.fromkeys(x for x in s if x))),
            fichiers=("fichier", lambda s: ", ".join(dict.fromkeys(s))),
            arrets=("dossier", "size"), dossiers=("dossier", "nunique"),
            debut=("ts_plan", "min"), fin=("ts_plan", "max")).reset_index()
        st.dataframe(det, hide_index=True, use_container_width=True, column_config={
            "tracteur": "Tracteur / ressource", "type_ressource": "Type", "societe": "Société affrétée",
            "chauffeurs": "Chauffeur(s)", "plateaux": "Plateau(x)",
            "fichiers": "Fichier(s)", "arrets": "Arrêts", "dossiers": "Dossiers",
            "debut": st.column_config.DatetimeColumn("Début", format="DD/MM/YY"),
            "fin": st.column_config.DatetimeColumn("Fin", format="DD/MM/YY")})

    # ── Calcul ──
    st.write("")
    if st.button("Calculer les km et la rentabilité", type="primary"):
        for key in ["rp_res"]:
            st.session_state.pop(key, None)
        bar, txt = st.progress(0.0), st.empty()

        def _p(msg, pct):
            txt.caption(msg)
            bar.progress(min(max(pct, 0.0), 1.0))
        try:
            res = calculer(sel, ca_dict, convention, reconstituer, progress_cb=_p)
            st.session_state["rp_res"] = dict(res=res, sel=sel, convention=convention,
                                              genere=pd.Timestamp.now(tz="Europe/Brussels").strftime("%d/%m/%Y %H:%M"))
            bar.progress(1.0)
            txt.empty()
        except Exception as e:
            st.session_state.pop("rp_res", None)
            st.error(f"Erreur pendant le calcul : {e}")

    if "rp_res" not in st.session_state:
        return
    df_traj, df_dos, df_flotte, df_ctrl = st.session_state["rp_res"]["res"]
    sel_c = st.session_state["rp_res"]["sel"]

    # ── Résultats ──
    ui.section("Résultats")
    conv = st.session_state["rp_res"].get("convention", "precedent")
    d0, d1 = sel_c["ts_plan"].min(), sel_c["ts_plan"].max()
    meta = dict(
        periode=(f"{d0:%d/%m/%Y} → {d1:%d/%m/%Y}" if pd.notna(d0) and pd.notna(d1) else ""),
        convention="Vide imputé au " + CONVENTIONS_VIDE.get(conv, conv).lower(),
        genere=st.session_state["rp_res"].get("genere", ""),
    )
    page_html = construire_tableau_de_bord(df_flotte, df_dos, df_traj, df_ctrl, sel_c, meta)

    b1, b2, b3, b4 = st.columns([1.15, 1.25, 1.1, 1], vertical_alignment="top")
    with b1:
        st.download_button("Rapport Excel", data=export_excel(df_flotte, df_dos, df_traj, df_ctrl),
                           file_name="Rapport_Renta_Plateau.xlsx", type="primary", use_container_width=True,
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with b2:
        st.download_button("Tableau de bord (HTML)", data=page_html.encode("utf-8"),
                           file_name="Renta_Plateau.html", mime="text/html", use_container_width=True)
    with b3:
        bouton_plein_ecran(page_html)
    with b4:
        if st.button("Nouveau calcul", use_container_width=True):
            st.session_state.pop("rp_res", None)
            st.rerun()
    components.html(page_html, height=1320, scrolling=True)


if os.environ.get("RENTA_PLATEAU_NO_UI") != "1":
    main()
