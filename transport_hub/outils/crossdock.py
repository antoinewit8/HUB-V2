"""
Page Streamlit : Plateforme cross-dock (version test)
Où placer un quai de groupage, et combien il rapporterait, d'après l'activité réelle.

Fichiers : LISTES_MISSIONS (arrêts par dossier) + export CA (prix par dossier),
joints par n° de dossier.
Calcul : tools/crossdock/engine.py (géolocalisation hors ligne du module Werbomont,
km = vol d'oiseau × 1,30).
"""

from __future__ import annotations

import hashlib
import io
import os
import sys
from datetime import datetime

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import pydeck as pdk
import streamlit as st

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import ui  # noqa: E402
from tools.crossdock import engine as X  # noqa: E402
from tools.werbomont import engine as W  # noqa: E402

ui.page("crossdock", "Où placer un quai de groupage, et ce qu’il rapporterait, d’après les dossiers réels.")

st.markdown("""
<style>
.cd-kpis { display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; margin: .4rem 0 1rem; }
.cd-kpi { background: #fff; border-radius: 18px; padding: 16px 18px; box-shadow: var(--ap-shadow); }
.cd-kpi .v { font-size: 30px; font-weight: 600; letter-spacing: -0.02em; line-height: 1.1; color: var(--ap-text); }
.cd-kpi .l { font-size: 13px; color: var(--ap-sub); margin-top: 4px; }
.cd-kpi .s { font-size: 12px; color: var(--ap-sub); margin-top: 2px; }
.cd-kpi.pos .v { color: var(--ap-green); } .cd-kpi.neg .v { color: var(--ap-red); }
@media (max-width: 760px) { .cd-kpis { grid-template-columns: repeat(2, 1fr); } }
</style>
""", unsafe_allow_html=True)

CARTO_LIGHT = "https://basemaps.cartocdn.com/gl/positron-gl-style/style.json"
# Sites : teintes foncées, distinctes du bleu (chargements) et de l'orange (déchargements)
COULEURS_SITES = [[29, 29, 31], [36, 138, 61], [137, 68, 171], [0, 125, 140], [99, 99, 102], [176, 66, 140]]


# ─── Utilitaires ─────────────────────────────────────────────────────────────
def nb(x) -> str:
    return "–" if x is None or pd.isna(x) else f"{x:_.0f}".replace("_", " ")


def eur(x) -> str:
    return "–" if x is None or pd.isna(x) else f"{x:_.0f} €".replace("_", " ")


def pct(x) -> str:
    return "–" if x is None or pd.isna(x) else f"{x:.0f} %"


def kpi(valeur, libelle, sous="", signe=None) -> str:
    cls = "" if signe is None else (" pos" if signe >= 0 else " neg")
    return (f'<div class="cd-kpi{cls}"><div class="v">{valeur}</div><div class="l">{libelle}</div>'
            f'<div class="s">{sous}</div></div>')


@st.cache_resource(show_spinner=False)
def geocodeur():
    return W.Geocodeur()


@st.cache_resource(show_spinner=False)
def lieux_proches():
    """Table des localités pour nommer le centre de gravité (pays voisins seulement)."""
    import gzip
    with gzip.open(W.CP_FILE, "rt", encoding="utf-8") as f:
        t = pd.read_csv(f, dtype=str)
    t = t[t["iso"].isin(["BE", "LU", "FR", "NL", "DE"])].copy()
    t["lat"], t["lon"] = t["lat"].astype(float), t["lon"].astype(float)
    return t.drop_duplicates(["iso", "place"]).reset_index(drop=True)


def nommer(lat, lon) -> tuple[str, str, str]:
    t = lieux_proches()
    i = int(np.argmin(W.hav(t["lat"].to_numpy(), t["lon"].to_numpy(), lat, lon)))
    r = t.iloc[i]
    return r["place"].title(), r["iso"], r["cp"].replace("L-", "")


@st.cache_data(show_spinner=False, max_entries=3)
def lire_excel(contenu: bytes) -> pd.DataFrame:
    return pd.read_excel(io.BytesIO(contenu), dtype={"Code postal": str, "C.P. chargement": str,
                                                     "C.P. déchargement": str, "N° Dossier": str})


@st.cache_data(show_spinner=False, max_entries=3)
def preparer(h_mis: str, h_ca: str, _mis: bytes, _ca: bytes | None):
    m = X.preparer_missions(lire_excel(_mis), geocodeur())
    ca = X.preparer_ca(lire_excel(_ca)) if _ca else None
    return X.construire_dossiers(m, ca), m["geo"].value_counts().to_dict()


# ─── Fichiers ────────────────────────────────────────────────────────────────
ui.section("Fichiers")
c1, c2 = st.columns(2)
with c1:
    f_mis = st.file_uploader("Liste des missions (LISTES_MISSIONS)", type=["xlsx", "xls"], key="cd_mis",
                             help="Une ligne par activité : CHARGER, DECHARGER… Donne les vrais arrêts, multi-points compris.")
with c2:
    f_ca = st.file_uploader("Chiffre d’affaires (export CA)", type=["xlsx", "xls"], key="cd_ca",
                            help="Une ligne par dossier avec le prix transport. Joint par n° de dossier.")

if not f_mis:
    st.markdown("""
<div class="ap-card"><h4>Ce que fait l’outil</h4>
<ul class="ap-list">
<li><b>Activité</b> <span>— chaque dossier est rebâti à partir de ses arrêts réels (missions), puis relié à son prix
(CA) par n° de dossier. On voit la part du groupage, en dossiers et en chiffre d’affaires.</span></li>
<li><b>Centre de gravité</b> <span>— le point qui minimise la distance aux chargements et aux déchargements,
pondéré par le nombre de dossiers ou le CA.</span></li>
<li><b>Sites candidats</b> <span>— Liège, Bettembourg, Metz, Werbomont, ou tout autre site saisi. Pour chacun :
détour de chaque dossier s’il passait par le quai.</span></li>
<li><b>Consolidation</b> <span>— les dossiers compatibles d’une même semaine vers une même zone partent ensemble
du quai. On compte les camions évités, les km gagnés et perdus.</span></li>
<li><b>Bilan</b> <span>— km nets × coût du km, moins la manutention et le coût du quai, ramené à l’année.</span></li>
</ul></div>
<div class="ap-note">Version test : km estimés au vol d’oiseau × 1,30, sans PTV. Les exports ne contiennent ni poids
ni mètres de plancher : le nombre de lots qu’un camion peut regrouper est une hypothèse à régler.
Le CA est facultatif ; sans lui, pas de pondération par CA.</div>""", unsafe_allow_html=True)
    st.stop()

b_mis = f_mis.getvalue()
b_ca = f_ca.getvalue() if f_ca else None
try:
    with st.spinner("Lecture, jointure et géolocalisation des dossiers…"):
        dossiers, precision_geo = preparer(hashlib.md5(b_mis).hexdigest(),
                                           hashlib.md5(b_ca).hexdigest() if b_ca else "",
                                           b_mis, b_ca)
except ValueError as e:
    st.error(str(e))
    st.stop()

if dossiers.empty or "type" not in dossiers:
    st.error("Aucun dossier avec un chargement et un déchargement dans la liste des missions.")
    st.stop()

# ─── Filtres ─────────────────────────────────────────────────────────────────
ok_all = dossiers[dossiers["statut"].eq("ok")]
dmin_, dmax_ = pd.to_datetime(ok_all["dt"]).min(), pd.to_datetime(ok_all["dt"]).max()
f1, f2, f3 = st.columns([2, 2, 1])
with f1:
    if pd.notna(dmin_):
        periode = st.date_input("Période (date du premier chargement)", value=(dmin_.date(), dmax_.date()),
                                min_value=dmin_.date(), max_value=dmax_.date(), format="DD/MM/YYYY")
    else:
        periode = None
with f2:
    types_tr = sorted(t for t in ok_all["type_transport"].fillna("").astype(str).unique() if t)
    sel_types = st.multiselect("Type de transport", types_tr, default=[], placeholder="Tous",
                               help="Le cross-dock vise surtout le plateau.")
with f3:
    excl_aff = st.toggle("Exclure les affrétés", value=True, help="Tracteur commençant par AFF.")

d = ok_all
if periode and isinstance(periode, tuple) and len(periode) == 2:
    dt = pd.to_datetime(d["dt"])
    d = d[(dt >= pd.Timestamp(periode[0])) & (dt < pd.Timestamp(periode[1]) + pd.Timedelta(days=1))]
if sel_types:
    d = d[d["type_transport"].isin(sel_types)]
if excl_aff:
    d = d[~d["tracteur"].fillna("").astype(str).str.upper().str.startswith("AFF")]
if d.empty:
    st.warning("Aucun dossier avec ces filtres.")
    st.stop()

jours = max((pd.to_datetime(d["dt"]).max() - pd.to_datetime(d["dt"]).min()).days + 1, 1)
mois = jours / 30.44
avec_ca = bool(d["dans_ca"].any())

# ─── Activité ────────────────────────────────────────────────────────────────
ui.section("Activité")
n = len(d)
n_grp = int(d["groupage"].sum())
ca_tot = d["prix"].sum() if avec_ca else np.nan
ca_grp = d.loc[d["groupage"], "prix"].sum() if avec_ca else np.nan
non_geo = int((dossiers["statut"] == "non géolocalisé").sum())
sans_ca = int((~d["dans_ca"]).sum())
st.markdown('<div class="cd-kpis">'
            + kpi(nb(n), "dossiers analysés", f"{nb(mois)} mois · {jours} jours")
            + kpi(pct(n_grp / n * 100), "en groupage", f"{nb(n_grp)} dossiers")
            + kpi(eur(ca_tot) if avec_ca else "–", "prix transport", "export CA" if avec_ca else "CA non fourni")
            + kpi(pct(ca_grp / ca_tot * 100) if avec_ca and ca_tot else "–", "du CA en groupage",
                  eur(ca_grp) if avec_ca else "")
            + "</div>", unsafe_allow_html=True)

T = d.groupby("type").agg(dossiers=("dossier", "size"), ca=("prix", "sum"), prix_moyen=("prix", "mean"),
                          clients=("client", "nunique")).reindex(
    ["Lot sec", "Groupage livraison", "Groupage chargement", "Groupage complet"]).dropna(how="all").reset_index()
T["part_dossiers"] = T["dossiers"] / n * 100
T["part_ca"] = T["ca"] / ca_tot * 100 if avec_ca and ca_tot else np.nan
st.dataframe(T[["type", "dossiers", "part_dossiers", "ca", "part_ca", "prix_moyen", "clients"]], hide_index=True,
             width="stretch", column_config={
                 "type": "Type", "dossiers": st.column_config.NumberColumn("Dossiers"),
                 "part_dossiers": st.column_config.NumberColumn("Part", format="%.1f %%"),
                 "ca": st.column_config.NumberColumn("Prix transport", format="%.0f €"),
                 "part_ca": st.column_config.NumberColumn("Part du CA", format="%.1f %%"),
                 "prix_moyen": st.column_config.NumberColumn("Prix moyen", format="%.0f €"),
                 "clients": st.column_config.NumberColumn("Clients")})
notes = []
if sans_ca and avec_ca:
    notes.append(f"{nb(sans_ca)} dossiers des missions sont absents du CA (prix inconnu, comptés à 0 €).")
if non_geo:
    notes.append(f"{nb(non_geo)} dossiers écartés : au moins un arrêt non géolocalisé.")
approx = sum(v for k, v in precision_geo.items() if k not in ("cp+localité", "manuel"))
if approx:
    notes.append(f"{nb(approx)} arrêts placés au code postal ou à la zone seulement (localité non reconnue).")
if notes:
    st.caption(" ".join(notes))

with st.expander("Groupage par client"):
    C = d.groupby("client").agg(dossiers=("dossier", "size"), groupage=("groupage", "sum"),
                                ca=("prix", "sum")).sort_values("dossiers", ascending=False)
    C["part_groupage"] = C["groupage"] / C["dossiers"] * 100
    st.dataframe(C.reset_index().head(60), hide_index=True, width="stretch", column_config={
        "client": "Client", "dossiers": "Dossiers", "groupage": "En groupage",
        "part_groupage": st.column_config.NumberColumn("Part groupage", format="%.0f %%"),
        "ca": st.column_config.NumberColumn("Prix transport", format="%.0f €")})

# ─── Centre de gravité ───────────────────────────────────────────────────────
ui.section("Centre de gravité")
g1, g2 = st.columns(2)
with g1:
    quoi = st.segmented_control("Points pris en compte", ["Tous", "Chargements", "Déchargements"], default="Tous",
                                key="cd_quoi") or "Tous"
with g2:
    choix_poids = ["Nombre de dossiers"] + (["Prix transport"] if avec_ca else [])
    poids_lib = st.segmented_control("Pondération", choix_poids, default="Nombre de dossiers", key="cd_poids") \
        or "Nombre de dossiers"
pts = X.points_flux(d, {"Tous": "tous", "Chargements": "chargements", "Déchargements": "dechargements"}[quoi],
                    "dossiers" if poids_lib == "Nombre de dossiers" else "ca")
cg, bary = X.centre_gravite(pts)
if cg is None:
    st.warning("Pas assez de points pour calculer un centre de gravité.")
    st.stop()
nom_cg, iso_cg, cp_cg = nommer(*cg)
st.markdown(f'<div class="ap-card"><h4>Près de {nom_cg} ({iso_cg} {cp_cg})</h4>'
            f'Point qui minimise la distance totale aux '
            f'{ {"Tous": "chargements et déchargements"}.get(quoi, quoi.lower())}, '
            f'pondérée par {"le nombre de dossiers" if poids_lib == "Nombre de dossiers" else "le prix transport"} : '
            f'{cg[0]:.4f}, {cg[1]:.4f}. Distance routière moyenne estimée aux points : '
            f'{nb(X.distance_moyenne(pts, cg))} km.<br><span style="color:var(--ap-sub);font-size:14px">'
            f'Le barycentre simple (moyenne des coordonnées) tombe près de {nommer(*bary)[0]} : il se laisse tirer '
            f'par les livraisons lointaines, c’est pourquoi on retient le point de distance minimale.</span></div>',
            unsafe_allow_html=True)

# ─── Sites candidats et hypothèses ───────────────────────────────────────────
ui.section("Sites candidats")
if "cd_sites" not in st.session_state:
    st.session_state["cd_sites"] = pd.DataFrame(X.SITES_DEFAUT, columns=["nom", "pays", "cp", "localite"])
sites_saisis = st.data_editor(st.session_state["cd_sites"], num_rows="dynamic", hide_index=True, width="stretch",
                              key="cd_sites_ed", column_config={
                                  "nom": st.column_config.TextColumn("Site", required=True),
                                  "pays": st.column_config.TextColumn("Pays", help="BE, LU, FR, NL, DE…"),
                                  "cp": st.column_config.TextColumn("Code postal"),
                                  "localite": st.column_config.TextColumn("Localité")})
ajouter_cg = st.toggle("Ajouter le centre de gravité comme site candidat", value=True)
sites = X.geolocaliser_sites(geocodeur(), sites_saisis)
introuvables = sites[sites["lat"].isna()]["nom"].tolist()
if introuvables:
    st.warning("Site(s) introuvable(s), vérifiez pays, code postal et localité : " + ", ".join(introuvables))
sites = sites[sites["lat"].notna()]
if ajouter_cg:
    sites = pd.concat([sites, pd.DataFrame([{"nom": f"Centre de gravité ({nom_cg})", "pays": iso_cg, "cp": cp_cg,
                                             "localite": nom_cg, "lat": cg[0], "lon": cg[1],
                                             "geo": "calculé"}])], ignore_index=True)
if sites.empty:
    st.stop()

with st.expander("Hypothèses de la simulation", expanded=True):
    h1, h2, h3 = st.columns(3)
    with h1:
        seuil_km = st.number_input("Détour max. par dossier (km)", 0, 500, 80, 10,
                                   help="Au-delà, le dossier reste en direct.")
        seuil_pct = st.number_input("Détour max. (% de la route)", 0, 200, 25, 5)
        groupage_seul = st.toggle("Seulement les dossiers en groupage", value=False,
                                  help="Sinon, les lots secs peuvent aussi être regroupés au quai.")
    with h2:
        lots_camion = st.number_input("Lots regroupés par camion", 1, 6, 2, 1,
                                      help="Hypothèse : combien de dossiers tiennent dans une remorque au départ du "
                                           "quai. Les exports ne donnent ni poids ni mètres de plancher. "
                                           "À 1, il n’y a aucune consolidation.")
        cout_km = st.number_input("Coût du km (€)", 0.0, 5.0, 1.20, 0.05, format="%.2f")
    with h3:
        manut_lot = st.number_input("Manutention par lot au quai (€)", 0, 1000, 40, 5,
                                    help="Déchargement + rechargement d’un lot sur le quai.")
        quai_mois = st.number_input("Coût fixe du quai par mois (€)", 0, 100_000, 3_000, 500,
                                    help="Location, personnel, chariot. 0 si un partenaire facture à l’acte.")
params = dict(seuil_km=float(seuil_km), seuil_pct=float(seuil_pct), lots_camion=int(lots_camion),
              groupage_seul=groupage_seul, cout_km=float(cout_km), manut_lot=float(manut_lot),
              quai_mois=float(quai_mois), mois=float(mois))

with st.spinner("Calcul des détours et de la consolidation…"):
    R, detail = X.comparer_sites(d, sites, pts, params)
R = R.sort_values("resultat_eur", ascending=False).reset_index(drop=True)

# ─── Comparaison ─────────────────────────────────────────────────────────────
ui.section("Comparaison des sites")
best = R.iloc[0]
st.markdown('<div class="cd-kpis">'
            + kpi(best["site"].split(" (")[0], "meilleur site", best["site"].split(" (")[1][:-1]
                  if " (" in best["site"] else "")
            + kpi(nb(best["camions_evites"]), "camions évités", f"{nb(best['consolides'])} dossiers regroupés")
            + kpi(nb(best["km_net"]) + " km", "km nets économisés",
                  f"{nb(best['km_economises'])} gagnés − {nb(best['km_ajoutes'])} ajoutés",
                  signe=best["km_net"])
            + kpi(eur(best["resultat_an_eur"]), "résultat estimé par an",
                  f"{eur(best['resultat_eur'])} sur la période", signe=best["resultat_an_eur"])
            + "</div>", unsafe_allow_html=True)
if int(lots_camion) == 1:
    st.markdown('<div class="ap-note">Avec 1 lot par camion, le quai ne regroupe rien : seuls les détours et les '
                'coûts restent. Montez à 2 ou 3 pour simuler une consolidation.</div>', unsafe_allow_html=True)

COLS_R = ["site", "compatibles", "consolides", "camions_evites", "km_economises", "km_ajoutes", "km_net",
          "gain_km_eur", "manutention_eur", "quai_eur", "resultat_eur", "resultat_an_eur", "detour_median",
          "dist_moy_flux"]
CFG_R = {
    "site": st.column_config.TextColumn("Site", width="medium"),
    "compatibles": st.column_config.NumberColumn("Dossiers compatibles", help="Détour sous les deux seuils."),
    "consolides": st.column_config.NumberColumn("Regroupés", help="Compatibles ET avec au moins un autre dossier "
                                                                  "la même semaine vers la même zone."),
    "camions_evites": st.column_config.NumberColumn("Camions évités"),
    "km_economises": st.column_config.NumberColumn("Km gagnés", format="%.0f"),
    "km_ajoutes": st.column_config.NumberColumn("Km ajoutés", format="%.0f", help="Détours + enchaînement des "
                                                                                    "livraisons dans la zone."),
    "km_net": st.column_config.NumberColumn("Km nets", format="%.0f"),
    "gain_km_eur": st.column_config.NumberColumn("Gain km", format="%.0f €"),
    "manutention_eur": st.column_config.NumberColumn("Manutention", format="%.0f €"),
    "quai_eur": st.column_config.NumberColumn("Quai", format="%.0f €"),
    "resultat_eur": st.column_config.NumberColumn("Résultat période", format="%.0f €"),
    "resultat_an_eur": st.column_config.NumberColumn("Résultat / an", format="%.0f €"),
    "detour_median": st.column_config.NumberColumn("Détour médian", format="%.0f km"),
    "dist_moy_flux": st.column_config.NumberColumn("Distance moy. aux flux", format="%.0f km"),
}
st.dataframe(R[COLS_R], hide_index=True, width="stretch", column_config=CFG_R)
B = R.iloc[::-1]
fig = go.Figure(go.Bar(
    x=B["resultat_an_eur"], y=B["site"].str.split(" \\(").str[0], orientation="h",
    marker_color=["#0071e3" if v >= 0 else "#d70015" for v in B["resultat_an_eur"]],
    text=[eur(v) for v in B["resultat_an_eur"]], textposition="outside", cliponaxis=False,
    hovertext=B["site"], hovertemplate="%{hovertext}<br>%{x:,.0f} € / an<extra></extra>"))
fig.update_layout(height=70 + 44 * len(B), margin=dict(l=10, r=90, t=10, b=30), plot_bgcolor="rgba(0,0,0,0)",
                  paper_bgcolor="rgba(0,0,0,0)", font=dict(family="Inter, -apple-system, Arial", size=13),
                  xaxis=dict(title="Résultat estimé par an (€)", gridcolor="#e5e5ea", zerolinecolor="#86868b",
                             separatethousands=True), bargap=.35, separators=", ")
st.plotly_chart(fig, config={"displayModeBar": False})

# ─── Carte ───────────────────────────────────────────────────────────────────
ui.section("Carte des flux")
agr = pts.assign(lat=pts["lat"].round(2), lon=pts["lon"].round(2)).groupby(["lat", "lon", "role"], as_index=False)["w"].sum()
agr["rayon"] = 2500 + np.sqrt(agr["w"] / agr["w"].max()) * 22000
agr["couleur"] = [[0, 113, 227, 150] if r == "C" else [255, 149, 0, 140] for r in agr["role"]]
agr["info"] = [("Chargement" if r == "C" else "Déchargement") + f" · poids {w:.1f}" for r, w in zip(agr["role"], agr["w"])]
S = R[["site", "lat", "lon", "resultat_an_eur", "camions_evites"]].copy()
S["couleur"] = [[215, 0, 21] if s.startswith("Centre de gravité") else COULEURS_SITES[i % len(COULEURS_SITES)][:3]
                for i, s in enumerate(S["site"])]
S["info"] = [f"{s} · {eur(v)} / an · {nb(c)} camions évités" for s, v, c in
             zip(S["site"], S["resultat_an_eur"], S["camions_evites"])]
S["etiquette"] = S["site"].str.split(" \\(").str[0]
couches = [
    pdk.Layer("ScatterplotLayer", data=agr, get_position="[lon, lat]", get_radius="rayon",
              get_fill_color="couleur", pickable=True, stroked=False),
    pdk.Layer("ScatterplotLayer", data=S, get_position="[lon, lat]", get_radius=7000, radius_min_pixels=7,
              get_fill_color="couleur", get_line_color=[255, 255, 255], line_width_min_pixels=2, stroked=True,
              pickable=True),
    pdk.Layer("TextLayer", data=S, get_position="[lon, lat]", get_text="etiquette", get_size=13,
              get_color=[29, 29, 31], get_pixel_offset=[0, -18], font_family='"Inter", Arial, sans-serif',
              font_weight=600, character_set="auto"),
]
st.pydeck_chart(pdk.Deck(layers=couches, map_style=CARTO_LIGHT, tooltip={"text": "{info}"},
                         initial_view_state=pdk.ViewState(latitude=float(cg[0]), longitude=float(cg[1]), zoom=5.2)),
                height=560)
st.caption("Bleu : chargements · orange : déchargements (taille selon la pondération choisie) · "
           "rouge : centre de gravité · points cerclés de blanc : sites candidats (survol pour le détail).")

# ─── Détail d’un site ────────────────────────────────────────────────────────
ui.section("Détail d’un site")
site_sel = st.selectbox("Site", R["site"].tolist(), index=0)
xd, gd = detail[site_sel]
t1, t2 = st.tabs(["Départs regroupés", "Dossiers"])
with t1:
    st.caption("Une ligne = les dossiers compatibles d’une même semaine vers une même zone, partis ensemble du quai.")
    st.dataframe(gd[["semaine", "zone_livraison", "dossiers", "camions", "camions_evites", "km_site_D1",
                     "km_economises", "km_detours", "km_tournee", "km_net", "ca", "liste"]],
                 hide_index=True, width="stretch", height=420, column_config={
                     "semaine": "Semaine", "zone_livraison": "Zone", "dossiers": "Dossiers", "camions": "Camions",
                     "camions_evites": "Évités",
                     "km_site_D1": st.column_config.NumberColumn("Quai → zone", format="%.0f km"),
                     "km_economises": st.column_config.NumberColumn("Km gagnés", format="%.0f"),
                     "km_detours": st.column_config.NumberColumn("Détours", format="%.0f km"),
                     "km_tournee": st.column_config.NumberColumn("Tournée", format="%.0f km"),
                     "km_net": st.column_config.NumberColumn("Km nets", format="%.0f"),
                     "ca": st.column_config.NumberColumn("Prix transport", format="%.0f €"),
                     "liste": st.column_config.TextColumn("Dossiers", width="large")})
with t2:
    seuls = st.toggle("Seulement les dossiers compatibles", value=True)
    xx = xd[xd["compatible"]] if seuls else xd
    xx = xx.assign(route=xx["ville_C1"] + " → " + xx["ville_Dn"])
    st.dataframe(xx[["dossier", "semaine", "client", "type", "route", "zone_livraison", "prix", "km_reel",
                     "km_via", "detour", "compatible", "consolide"]].sort_values("detour"),
                 hide_index=True, width="stretch", height=420, column_config={
                     "dossier": "Dossier", "semaine": "Semaine", "client": "Client", "type": "Type",
                     "route": st.column_config.TextColumn("Route", width="large"), "zone_livraison": "Zone",
                     "prix": st.column_config.NumberColumn("Prix", format="%.0f €"),
                     "km_reel": st.column_config.NumberColumn("Km route", format="%.0f"),
                     "km_via": st.column_config.NumberColumn("Km via quai", format="%.0f"),
                     "detour": st.column_config.NumberColumn("Détour", format="%.0f km"),
                     "compatible": "Compatible", "consolide": "Regroupé"})


# ─── Export Excel ────────────────────────────────────────────────────────────
def export_excel() -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    per = (f"{pd.Timestamp(periode[0]):%d/%m/%Y} – {pd.Timestamp(periode[1]):%d/%m/%Y}"
           if periode and isinstance(periode, tuple) and len(periode) == 2 else "")
    synth = pd.DataFrame([
        ("Période", per), ("Dossiers analysés", n), ("Part groupage (dossiers)", f"{n_grp / n * 100:.1f} %"),
        ("Part groupage (CA)", f"{ca_grp / ca_tot * 100:.1f} %" if avec_ca and ca_tot else "CA non fourni"),
        ("Centre de gravité", f"près de {nom_cg} ({iso_cg} {cp_cg}) · {cg[0]:.4f}, {cg[1]:.4f} · {quoi}, {poids_lib}"),
        ("Détour max.", f"{seuil_km} km et {seuil_pct} % de la route"),
        ("Lots par camion (hypothèse)", lots_camion), ("Coût du km", f"{cout_km:.2f} €"),
        ("Manutention par lot", f"{manut_lot} €"), ("Quai par mois", f"{quai_mois} €"),
        ("Km", "vol d'oiseau × 1,30 (sans PTV)"),
        ("Généré le", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ], columns=["Élément", "Valeur"])
    sites_x = R[COLS_R].rename(columns={k: (CFG_R[k].get("label") if isinstance(CFG_R[k], dict) else None) or k
                                        for k in COLS_R})
    dos_x = xd.assign(route=xd["ville_C1"] + " → " + xd["ville_Dn"])[
        ["dossier", "semaine", "client", "type", "route", "zone_livraison", "prix", "km_reel", "km_via", "detour",
         "compatible", "consolide"]]
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        synth.to_excel(xw, sheet_name="Synthèse", index=False)
        sites_x.to_excel(xw, sheet_name="Sites", index=False)
        T.to_excel(xw, sheet_name="Activité", index=False)
        gd.to_excel(xw, sheet_name="Départs regroupés", index=False)
        dos_x.to_excel(xw, sheet_name="Dossiers", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(color="FFFFFF", bold=True)
                c.fill = PatternFill("solid", fgColor="003087")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "B2"
            ws.row_dimensions[1].height = 30
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), 1):
                larg = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, larg + 2), 60)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "#,##0"
            ws.auto_filter.ref = ws.dimensions
        xw.book["Synthèse"].column_dimensions["B"].width = 90
    return out.getvalue()


st.download_button("Télécharger l’étude (Excel)", export_excel(), f"crossdock_{datetime.now():%Y%m%d}.xlsx",
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
st.caption("Les onglets « Départs regroupés » et « Dossiers » portent sur le site choisi dans le détail.")
